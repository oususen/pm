"""社内AI用のMarkdownナレッジ検索。

ベクトルDBへ全文を保存せず、リポジトリの正式原本を都度検索する。
設定画面のAIKnowledgeSourceで有効になっている区分だけを対象にする。
"""
from __future__ import annotations

import re
import csv
import os
import subprocess
from functools import lru_cache
from pathlib import Path

from ai.config.models import AIKnowledgeDocument, AIKnowledgeSource


PROJECT_ROOT = Path(__file__).resolve().parents[4]
MANUAL_ROOT = PROJECT_ROOT / 'pm-ui' / 'public' / 'manual'
SPEC_ROOT = PROJECT_ROOT / '仕様書'
LOCAL_KNOWLEDGE_ROOT = PROJECT_ROOT / 'pm_backend' / 'apps' / 'ai' / 'knowledge'
MAX_CHUNK_CHARS = 1400
OCR_EXECUTABLE = os.environ.get('OCR_TESSERACT_PATH', r'C:\Program Files\Tesseract-OCR\tesseract.exe')
OCR_TESSDATA_DIR = os.environ.get(
    'OCR_TESSDATA_DIR', str(PROJECT_ROOT / 'pm_backend' / 'ocr_data'),
)
IMAGE_SUFFIXES = {'.png', '.jpg', '.jpeg', '.webp', '.bmp'}
SCREEN_KEYWORDS = {
    'orders': ('受注', '注文', 'ルーティング'),
    'production': ('生産', '工程', '実績', '作業'),
    'quality': ('品質', '仕損', '不良', '検査'),
    'overtime': ('勤務', '残業', '申請', '勤怠'),
    'purchase': ('仕入', '購買', '発注', '検収'),
    'shipping': ('出荷', '納入', '便'),
    'inventory': ('在庫', '棚卸', '進度'),
}


def _enabled_categories():
    """DB設定を読み、未登録時は何も検索しない。"""
    return set(
        AIKnowledgeSource.objects.filter(is_enabled=True).values_list('category', flat=True)
    )


def _paths_for_categories(categories, document_ids=None):
    """区分に対応する正式原本だけを返す。READMEは案内文のため検索しない。"""
    paths = []
    document_ids = tuple(document_ids or ())
    # 資料を指定した会話では、その資料だけを検索し、応答を速く・根拠を明確にする。
    # 明示添付は利用者自身の今回限りの指定なので、ナレッジ一覧の有効/無効には従わない。
    # 添付しない通常会話だけは、下段の is_enabled=True で自動検索対象を限定する。
    if document_ids:
        documents = AIKnowledgeDocument.objects.filter(
            category__in=categories, id__in=document_ids,
        )
        for document in documents:
            try:
                path = Path(document.file.path)
            except (ValueError, OSError):
                continue
            if path.exists():
                paths.append((path, f'AIナレッジ資料/{document.name}'))
        return sorted(set(paths), key=lambda item: item[1])
    if 'manual' in categories and MANUAL_ROOT.exists():
        paths.extend((path, path.relative_to(PROJECT_ROOT).as_posix()) for path in MANUAL_ROOT.rglob('*.md') if path.name.lower() != 'readme.md')
    if 'pm_structure' in categories:
        for name in ('PMアプリ構造辞書.md', '社内AIチャット仕様書.md'):
            path = SPEC_ROOT / name
            if path.exists():
                paths.append((path, path.relative_to(PROJECT_ROOT).as_posix()))
    if 'security' in categories:
        path = SPEC_ROOT / '社内AI運用規約・AI用DB辞書.md'
        if path.exists():
            paths.append((path, path.relative_to(PROJECT_ROOT).as_posix()))
    if 'procedure' in categories:
        procedure_root = LOCAL_KNOWLEDGE_ROOT / 'procedures'
        if procedure_root.exists():
            paths.extend((path, path.relative_to(PROJECT_ROOT).as_posix()) for path in procedure_root.rglob('*.md') if path.name.lower() != 'readme.md')
        # 承認済みの運用手順・トラブル対応は既存仕様書を正式原本として参照する。
        paths.extend(
            (path, path.relative_to(PROJECT_ROOT).as_posix())
            for path in SPEC_ROOT.glob('*.md')
            if any(word in path.name for word in ('手順', '対応'))
        )
    for document in AIKnowledgeDocument.objects.filter(is_enabled=True, category__in=categories):
        try:
            path = Path(document.file.path)
        except (ValueError, OSError):
            continue
        if path.exists():
            paths.append((path, f'AIナレッジ資料/{document.name}'))
    return sorted(set(paths), key=lambda item: item[1])


def _normalize(value):
    return re.sub(r'[^0-9a-zA-Zぁ-んァ-ヶ一-龠]', '', str(value).lower())


def _keywords(value):
    normalized = _normalize(value)
    terms = set(re.findall(r'[a-z0-9]{2,}', normalized))
    # 日本語は形態素解析に依存せず、2文字単位の一致で候補を絞る。
    terms.update(normalized[index:index + 2] for index in range(len(normalized) - 1))
    return {term for term in terms if term}


def _split_sections(text):
    sections = []
    heading = '本文'
    buffer = []
    for line in text.splitlines():
        if line.startswith('#'):
            if buffer:
                sections.append((heading, '\n'.join(buffer).strip()))
                buffer = []
            heading = line.lstrip('#').strip() or '本文'
        buffer.append(line)
    if buffer:
        sections.append((heading, '\n'.join(buffer).strip()))

    chunks = []
    for section_heading, body in sections:
        for start in range(0, max(len(body), 1), MAX_CHUNK_CHARS):
            chunk = body[start:start + MAX_CHUNK_CHARS].strip()
            if chunk:
                chunks.append((section_heading, chunk))
    return chunks


def _extract_spreadsheet_text(path):
    """表計算資料は見出しと値をテキスト化する。式は再計算せず保存済みの値だけを読む。"""
    suffix = path.suffix.lower()
    if suffix in {'.csv', '.tsv'}:
        delimiter = '\t' if suffix == '.tsv' else ','
        for encoding in ('utf-8-sig', 'cp932'):
            try:
                with path.open('r', encoding=encoding, newline='') as stream:
                    rows = list(csv.reader(stream, delimiter=delimiter))[:300]
                return '\n'.join(' | '.join(cell.strip() for cell in row[:50] if cell.strip()) for row in rows)
            except UnicodeDecodeError:
                continue
        return ''
    if suffix == '.xlsx':
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=True)
        rows = []
        for sheet in workbook.worksheets[:20]:
            rows.append(f'## シート: {sheet.title}')
            for row in sheet.iter_rows(max_row=300, max_col=50, values_only=True):
                values = [str(value).strip() for value in row if value not in (None, '')]
                if values:
                    rows.append(' | '.join(values))
        workbook.close()
        return '\n'.join(rows)
    if suffix == '.xls':
        import pandas as pd
        rows = []
        workbook = pd.ExcelFile(path, engine='xlrd')
        try:
            for sheet_name in workbook.sheet_names[:20]:
                rows.append(f'## シート: {sheet_name}')
                frame = pd.read_excel(workbook, sheet_name=sheet_name, header=None, nrows=300).iloc[:, :50]
                for row in frame.fillna('').astype(str).values.tolist():
                    values = [value.strip() for value in row if value.strip()]
                    if values:
                        rows.append(' | '.join(values))
        finally:
            workbook.close()
        return '\n'.join(rows)
    return ''


def _extract_text(path):
    suffix = path.suffix.lower()
    if suffix == '.pdf':
        from pypdf import PdfReader
        reader = PdfReader(path)
        return '\n'.join((page.extract_text() or '') for page in reader.pages[:100])
    if suffix in {'.xlsx', '.xls', '.csv', '.tsv'}:
        return _extract_spreadsheet_text(path)
    if suffix in IMAGE_SUFFIXES:
        return _extract_image_text(path)
    return path.read_text(encoding='utf-8')


def _extract_image_text(path):
    """ローカルTesseractで画像内の日本語・英数字をOCRする。画像は外部へ送信しない。"""
    executable = Path(OCR_EXECUTABLE)
    tessdata_dir = Path(OCR_TESSDATA_DIR)
    if not executable.exists() or not tessdata_dir.exists():
        return ''
    result = subprocess.run(
        [
            str(executable), str(path), 'stdout', '-l', 'jpn+eng',
            '--tessdata-dir', str(tessdata_dir), '--psm', '6',
        ],
        capture_output=True,
        text=True,
        encoding='utf-8',
        timeout=30,
        check=False,
    )
    return result.stdout.strip() if result.returncode == 0 else ''


@lru_cache(maxsize=1024)
def _read_chunks(path_text, modified_ns):
    """更新時刻をキャッシュキーに含め、開発中のMarkdown編集を次回検索へ反映する。"""
    path = Path(path_text)
    try:
        return tuple(_split_sections(_extract_text(path)))
    # 利用者が登録した壊れた資料1件で、AIチャット全体を失敗させない。
    except Exception:  # 解析対象は利用者が登録する外部ファイルのため、形式不正も無視する。
        return ()


def retrieve_knowledge(question, screen_context=None, limit=4, max_chars=5200, document_ids=None):
    """質問と起点画面に近い正式原本の断片を返す。"""
    categories = _enabled_categories()
    if not categories:
        return []
    query_terms = _keywords(question)
    screen_id = (screen_context or {}).get('id', 'ai_home')
    screen_terms = set(SCREEN_KEYWORDS.get(screen_id, ()))
    candidates = []
    for path, relative in _paths_for_categories(categories, document_ids):
        try:
            modified_ns = path.stat().st_mtime_ns
        except OSError:
            continue
        path_terms = _keywords(relative)
        for heading, content in _read_chunks(str(path), modified_ns):
            content_terms = _keywords(content)
            score = len(query_terms & content_terms) * 4 + len(query_terms & path_terms) * 7
            # 利用者がAI設定から追加した資料は、その固有の記述を優先して参照する。
            if relative.startswith('AIナレッジ資料/'):
                score += len(query_terms & content_terms) * 8 + 12
            # 起点画面の用語が原本のパスまたは本文にあれば優先する。
            score += sum(5 for term in screen_terms if term in relative or term in content)
            if score <= 0:
                continue
            candidates.append({
                'path': relative,
                'heading': heading,
                'content': content,
                'score': score,
            })
    selected = []
    used_chars = 0
    per_path_count = {}
    for item in sorted(candidates, key=lambda value: (-value['score'], value['path'], value['heading'])):
        if (
            len(selected) >= limit
            or used_chars + len(item['content']) > max_chars
            or per_path_count.get(item['path'], 0) >= 2
        ):
            continue
        selected.append({key: item[key] for key in ('path', 'heading', 'content')})
        used_chars += len(item['content'])
        per_path_count[item['path']] = per_path_count.get(item['path'], 0) + 1
    return selected


def knowledge_prompt(chunks):
    """LLMへ渡す根拠付きコンテキストを作る。"""
    if not chunks:
        return ''
    sections = []
    for chunk in chunks:
        sections.append(
            f"[出典: {chunk['path']} / {chunk['heading']}]\n{chunk['content']}"
        )
    return (
        '\n\n以下は社内の正式ナレッジ原本から検索した参考情報です。'
        'この内容に根拠がある場合だけ回答に使い、根拠がないことは推測しないでください。'
        'DBの数値が必要な質問は、必ずDBツールまたはDB集計結果を優先してください。\n\n'
        + '\n\n---\n\n'.join(sections)
    )


def knowledge_source_label(chunks):
    """画面の根拠データ欄に表示する短い出典名。"""
    paths = []
    for chunk in chunks:
        if chunk['path'] not in paths:
            paths.append(chunk['path'])
    return 'ナレッジ: ' + '、'.join(paths[:2]) if paths else ''
