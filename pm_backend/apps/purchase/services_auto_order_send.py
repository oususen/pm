import math
from collections import defaultdict
from datetime import date, timedelta
from decimal import Decimal, ROUND_HALF_UP
from io import BytesIO

from django.db import transaction
from django.db.models import Q, Sum
from openpyxl import Workbook

from masters.models import CalendarDay, Product, Supplier
from orders.utils.calendar_utils import WorkingDayCalculator
from production.models import LineDemand
from production.models_line_backlog import LineBacklog

from .order_proposal_views import (
    _generate_raw_pattern_dates,
    _get_daiso_calendar,
    _shift_to_previous_working_day,
)
from .process_resolver import resolve_purchase_line, resolve_supplier_process
from .tasks_auto_delivery_list import _recalculate_supplier_progress_for_auto_delivery
from .views import _sort_delivery_list_items


def _resolve_planning_horizon_end(calc, base_date, progress_days_forward):
    return calc.add_working_days(base_date, int(progress_days_forward or 30))


def _resolve_shifted_pattern_dates(schedule, window_start, window_end, calc):
    shifted_dates = set()
    for raw_date in _generate_raw_pattern_dates(schedule, window_start, window_end, calc):
        shifted_date = _shift_to_previous_working_day(raw_date, calc)
        if window_start <= shifted_date <= window_end:
            shifted_dates.add(shifted_date)
    return sorted(shifted_dates)


def _resolve_supplier_calendar_delivery_dates(supplier, window_start, window_end):
    calendar_id = getattr(supplier, 'calendar_id', None)
    if not calendar_id:
        return []
    return list(
        CalendarDay.objects.filter(
            calendar_id=calendar_id,
            target_date__gte=window_start,
            target_date__lte=window_end,
            is_delivery_day=True,
        )
        .order_by('target_date')
        .values_list('target_date', flat=True)
    )


def resolve_delivery_cycles(config, base_date=None):
    from .models import SupplierOrderSchedule

    supplier = config.supplier
    today = base_date or date.today()
    daiso_cal = _get_daiso_calendar()
    calc = WorkingDayCalculator(daiso_cal)
    lead_time_days = 5 if config.lead_time_days is None else int(config.lead_time_days)
    earliest_delivery_date = calc.add_working_days(today, lead_time_days)
    delivery_day_mode = getattr(config, 'delivery_day_mode', 'PATTERN')

    schedule = SupplierOrderSchedule.objects.filter(
        supplier_id=supplier.id,
        is_enabled=True,
    ).select_related('pattern').first()
    if delivery_day_mode == 'PATTERN' and (not schedule or not schedule.pattern):
        return {
            'status': 'SKIPPED',
            'message': f'仕入先 {supplier.supplier_code} の納入パターンが未設定です',
            'cycles': [],
            'schedule': schedule,
            'first_delivery_date': earliest_delivery_date,
        }

    horizon_end = _resolve_planning_horizon_end(calc, today, config.progress_days_forward)
    window_start = earliest_delivery_date - timedelta(days=7)
    window_end = horizon_end + timedelta(days=120)
    if delivery_day_mode == 'SUPPLIER_CALENDAR':
        if not getattr(supplier, 'calendar_id', None):
            return {
                'status': 'SKIPPED',
                'message': f'仕入先 {supplier.supplier_code} の仕入れ先カレンダが未設定です',
                'cycles': [],
                'schedule': schedule,
                'first_delivery_date': earliest_delivery_date,
            }
        pattern_dates = _resolve_supplier_calendar_delivery_dates(supplier, window_start, window_end)
    else:
        pattern_dates = _resolve_shifted_pattern_dates(schedule, window_start, window_end, calc)

    if earliest_delivery_date not in pattern_dates:
        return {
            'status': 'SKIPPED',
            'message': f'{earliest_delivery_date} は {supplier.supplier_code} の納入日ではありません',
            'cycles': [],
            'schedule': schedule,
            'first_delivery_date': earliest_delivery_date,
        }

    first_delivery_date = earliest_delivery_date
    next_delivery_date = next((d for d in pattern_dates if d > first_delivery_date), None)
    coverage_dates = []
    if next_delivery_date:
        cursor = first_delivery_date
        while cursor < next_delivery_date:
            coverage_dates.append(cursor)
            cursor += timedelta(days=1)
    else:
        coverage_dates = [first_delivery_date]

    cycles = [{
        'delivery_date': first_delivery_date,
        'next_delivery_date': next_delivery_date,
        'coverage_dates': coverage_dates,
    }]

    return {
        'status': 'SUCCESS',
        'message': '',
        'cycles': cycles,
        'schedule': schedule,
        'first_delivery_date': first_delivery_date,
    }


def _collect_supplier_product_ids(supplier):
    from masters.models import BOMItem, RoutingStep

    bom_product_ids = set(
        BOMItem.objects.filter(
            supplier_id=supplier.id,
            child_product_id__isnull=False,
        ).values_list('child_product_id', flat=True).distinct()
    )
    routing_product_ids = set(
        RoutingStep.objects.filter(
            supplier_id=supplier.id,
            output_product_id__isnull=False,
        ).values_list('output_product_id', flat=True).distinct()
    )
    return sorted(bom_product_ids | routing_product_ids)


def _load_first_cycle_progress_basis(line, product_ids, delivery_date, send_date):
    """初回納入回の基準値を返す。

    納入日当日の需要を coverage に含めるため、
    基準計進は納入日前日の値を採用する。

    計画進度は業務日の前日以前には実績を使用しているため、
    実績と計画の差分を補正するのは業務日当日以降だけにする。

    ルール:
    - 基準日は納入日前日
    - 基準日が業務日当日以降かつ実績あり: 計進 + 実績 − 計画
    - 基準日と異なる業務日当日に実績あり: さらに業務日当日の (実績 − 計画) を加算
    """
    progress_map = {}
    if not product_ids:
        return progress_map
    basis_date = delivery_date - timedelta(days=1)

    for product_id in product_ids:
        day_qs = LineBacklog.objects.filter(
            line_id=line.id,
            product_id=product_id,
            plan_date=basis_date,
        )
        planned_progress = int(day_qs.aggregate(v=Sum('planned_progress_qty'))['v'] or 0)
        actual_qty = int(day_qs.filter(sequence_no=0).aggregate(v=Sum('actual_qty'))['v'] or 0)
        plan_qty = int(day_qs.filter(sequence_no=1).aggregate(v=Sum('plan_qty'))['v'] or 0)

        basis = planned_progress
        # 過去日の planned_progress_qty は進度再計算で既に actual_qty を採用している。
        if actual_qty > 0 and basis_date >= send_date:
            basis = planned_progress + actual_qty - plan_qty

        # 送信日の実績補正: 計進は送信日を plan で計算するため、実績との差を反映
        if send_date != basis_date:
            today_qs = LineBacklog.objects.filter(
                line_id=line.id,
                product_id=product_id,
                plan_date=send_date,
            )
            today_actual = int(today_qs.filter(sequence_no=0).aggregate(v=Sum('actual_qty'))['v'] or 0)
            if today_actual > 0:
                today_plan = int(today_qs.filter(sequence_no=1).aggregate(v=Sum('plan_qty'))['v'] or 0)
                basis += today_actual - today_plan

        progress_map[product_id] = basis

    return progress_map


def _round_order_qty(required_qty, product, calc_mode):
    required_qty = max(0, int(required_qty or 0))
    if required_qty <= 0:
        return 0
    if calc_mode != 'LOT_ROUNDED':
        return required_qty

    multiple = int(getattr(product, 'order_lot_multiple', 0) or 1)
    if multiple <= 0:
        multiple = 1
    min_lot = int(getattr(product, 'order_lot_min', 0) or 0)
    base_qty = max(required_qty, min_lot)
    rounded_qty = int(math.ceil(base_qty / multiple) * multiple)
    return rounded_qty


def _build_cycle_product_map(line, coverage_dates):
    product_map = {}
    demand_qs = (
        LineDemand.objects.filter(line=line, plan_date__in=coverage_dates)
        .filter(Q(firm_qty__gt=0) | Q(forecast_qty__gt=0))
        .select_related('product')
    )
    for demand in demand_qs:
        product = demand.product
        if not product:
            continue
        product_id = product.id
        if product_id not in product_map:
            product_map[product_id] = {
                'product': product,
                'product_id': product_id,
                'product_code': product.product_code,
                'product_name': product.product_name,
                'transfer_destination': product.transfer_destination or '',
                'transfer_destination_label': product.get_transfer_destination_display() if product.transfer_destination else '',
                'coverage_demand': 0,
                'daily': {},
            }
        qty = int(demand.firm_qty or 0) + int(demand.forecast_qty or 0)
        product_map[product_id]['coverage_demand'] += qty
        day_key = demand.plan_date.isoformat()
        product_map[product_id]['daily'][day_key] = product_map[product_id]['daily'].get(day_key, 0) + qty
    return product_map


def _load_safety_stock_map(product_ids):
    from production.models_production import StockAllocation
    if not product_ids:
        return {}
    return dict(
        StockAllocation.objects.filter(product_id__in=product_ids)
        .values_list('product_id', 'min_stock_qty')
    )


@transaction.atomic
def simulate_and_save_order_plans(config, base_date=None):
    supplier = config.supplier
    line = resolve_purchase_line(supplier)
    if not line:
        return {
            'status': 'FAILED',
            'message': f'仕入先 {supplier.supplier_code} に対応するラインが見つかりません',
            'cycles': [],
            'items': [],
        }

    cycle_result = resolve_delivery_cycles(config, base_date=base_date)
    if cycle_result['status'] != 'SUCCESS':
        return {
            'status': cycle_result['status'],
            'message': cycle_result['message'],
            'cycles': [],
            'items': [],
        }

    cycles = cycle_result['cycles']
    if not cycles:
        return {
            'status': 'SUCCESS',
            'message': cycle_result['message'],
            'cycles': [],
            'items': [],
        }

    planning_horizon_end = _resolve_planning_horizon_end(
        WorkingDayCalculator(_get_daiso_calendar()),
        base_date or date.today(),
        config.progress_days_forward,
    )
    initial_days_forward = max((planning_horizon_end - (base_date or date.today())).days + 1, 7)
    product_ids = _collect_supplier_product_ids(supplier)
    _recalculate_supplier_progress_for_auto_delivery(
        supplier,
        line,
        days_back=max(int(config.progress_days_back or 7), 1),
        days_forward=initial_days_forward,
        product_ids=product_ids,
    )

    safety_stock_enabled = getattr(config, 'safety_stock_enabled', False)
    safety_stock_multiplier = float(getattr(config, 'safety_stock_multiplier', 1) or 1)
    safety_stock_map = _load_safety_stock_map(product_ids) if safety_stock_enabled else {}

    first_delivery_date = cycles[0]['delivery_date']
    today = base_date or date.today()
    current_progress_map = _load_first_cycle_progress_basis(line, product_ids, first_delivery_date, send_date=today)
    all_items = []
    recalculated_product_ids = set()

    for cycle in cycles:
        delivery_date = cycle['delivery_date']
        coverage_dates = cycle['coverage_dates']
        product_map = _build_cycle_product_map(line, coverage_dates)
        for info in product_map.values():
            product = info['product']
            product_id = info['product_id']
            coverage_demand = int(info['coverage_demand'])
            current_planned_progress = int(current_progress_map.get(product_id, 0))
            safety_addition = 0
            if safety_stock_enabled:
                min_stock = int(safety_stock_map.get(product_id, 0) or 0)
                safety_addition = int(Decimal(str(min_stock * safety_stock_multiplier)).quantize(Decimal('1'), rounding=ROUND_HALF_UP))
            required_qty = max(0, coverage_demand - current_planned_progress + safety_addition)
            process = resolve_supplier_process(
                supplier=supplier,
                line=line,
                product=product,
                create_purchase_process=True,
            )
            if not process:
                continue

            if required_qty <= 0:
                LineBacklog.objects.filter(
                    line_id=line.id,
                    process_id=process.id,
                    product_id=product_id,
                    plan_date=delivery_date,
                    sequence_no=1,
                ).delete()
                continue

            plan_qty = _round_order_qty(required_qty, product, config.calc_mode)

            LineBacklog.objects.update_or_create(
                line_id=line.id,
                process_id=process.id,
                product_id=product_id,
                plan_date=delivery_date,
                sequence_no=1,
                defaults={
                    'plan_qty': plan_qty,
                    'actual_qty': 0,
                },
            )
            current_progress_map[product_id] = current_planned_progress + plan_qty - coverage_demand
            recalculated_product_ids.add(product_id)
            all_items.append({
                'product_id': product_id,
                'product_code': info['product_code'],
                'product_name': info['product_name'],
                'transfer_destination': info['transfer_destination'],
                'transfer_destination_label': info['transfer_destination_label'],
                'delivery_date': delivery_date,
                'expected_qty': plan_qty,
                'coverage_demand': coverage_demand,
                'required_qty': required_qty,
                'daily': info['daily'],
            })

    if recalculated_product_ids:
        horizon_end = max(cycle['coverage_dates'][-1] for cycle in cycles if cycle['coverage_dates'])
        _recalculate_supplier_progress_for_auto_delivery(
            supplier,
            line,
            days_back=max(int(config.progress_days_back or 7), 1),
            days_forward=max((horizon_end - (base_date or date.today())).days + 1, 7),
            product_ids=sorted(recalculated_product_ids),
        )

    return {
        'status': 'SUCCESS',
        'message': f'{len(all_items)}件の注文計画を保存しました',
        'cycles': cycles,
        'items': all_items,
        'line': line,
    }


def generate_order_excel(items, supplier):
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = '注文書'

    delivery_dates = sorted({item['delivery_date'] for item in items})
    coverage_dates = sorted({date.fromisoformat(day) for item in items for day in item.get('daily', {}).keys()})
    day_names = ['月', '火', '水', '木', '金', '土', '日']
    date_headers = [f'{d.month}/{d.day}({day_names[d.weekday()]})' for d in coverage_dates]

    headers = ['品番', '品名', '移動先', '数量', '納品日'] + date_headers + ['仕入先コード', '伝票番号']
    ws.append(headers)

    header_fill = PatternFill(start_color='1D4ED8', end_color='1D4ED8', fill_type='solid')
    header_font = Font(color='FFFFFF', bold=True, size=10)
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin'),
    )

    sorted_items = sorted(
        _sort_delivery_list_items(items),
        key=lambda item: (item['delivery_date'], item['product_code']),
    )
    for item in sorted_items:
        row = [
            item['product_code'],
            item['product_name'],
            item.get('transfer_destination_label', ''),
            int(item.get('expected_qty') or 0),
            item['delivery_date'].isoformat(),
        ]
        for coverage_date in coverage_dates:
            row.append(item.get('daily', {}).get(coverage_date.isoformat(), ''))
        row += [supplier.supplier_code or '', '']
        ws.append(row)

    for col_idx, _header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = thin_border

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.border = thin_border
            if cell.column == 4:
                cell.alignment = Alignment(horizontal='right')
            else:
                cell.alignment = Alignment(horizontal='left')

    widths = {
        1: 18,
        2: 28,
        3: 12,
        4: 10,
        5: 12,
    }
    for idx in range(6, 6 + len(date_headers)):
        widths[idx] = 11
    widths[len(headers) - 1] = 14
    widths[len(headers)] = 14
    for col_idx, width in widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions

    note = wb.create_sheet('説明')
    note.append(['【注文書の見方】'])
    note.append(['', '・「数量」は今回保存した計画数です。'])
    note.append(['', '・日別内訳は元需要です。数量合計と一致しない場合があります。'])
    note.append(['', f'・対象納入回数: {len(delivery_dates)}回'])
    note.append(['', f'・対象納入日: {", ".join(d.isoformat() for d in delivery_dates)}'])

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def generate_order_pdf(items, supplier):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Spacer, Paragraph
    from reportlab.lib.styles import ParagraphStyle
    from shipping.services.shipping_pdf_generator import register_japanese_fonts

    register_japanese_fonts()
    font_name = 'MSGothic'

    sorted_items = sorted(
        items,
        key=lambda item: (item['delivery_date'], item['product_code']),
    )

    buf = BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=landscape(A4),
        leftMargin=15 * mm,
        rightMargin=15 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
    )

    elements = []
    title_style = ParagraphStyle('Title', fontName=font_name, fontSize=14, leading=18)
    sub_style = ParagraphStyle('Sub', fontName=font_name, fontSize=9, leading=12)

    delivery_dates = sorted({item['delivery_date'] for item in items})
    elements.append(Paragraph(f'注文書 — {supplier.supplier_name}', title_style))
    elements.append(Spacer(1, 3 * mm))
    elements.append(Paragraph(
        f'対象納入日: {", ".join(d.isoformat() for d in delivery_dates)}　／　{len(items)}件',
        sub_style,
    ))
    elements.append(Spacer(1, 5 * mm))

    headers = ['品番', '品名', '移動先', '数量', '納品日', '伝票番号']
    cell_style = ParagraphStyle('Cell', fontName=font_name, fontSize=9, leading=11)
    header_style = ParagraphStyle('Header', fontName=font_name, fontSize=9, leading=11, textColor=colors.white)

    data = [[Paragraph(h, header_style) for h in headers]]
    for item in sorted_items:
        qty = int(item.get('expected_qty') or 0)
        data.append([
            Paragraph(item['product_code'], cell_style),
            Paragraph(item['product_name'], cell_style),
            Paragraph(item.get('transfer_destination_label', ''), cell_style),
            Paragraph(str(qty), cell_style),
            Paragraph(item['delivery_date'].isoformat(), cell_style),
            Paragraph('', cell_style),
        ])

    col_widths = [45 * mm, 70 * mm, 30 * mm, 25 * mm, 30 * mm, 35 * mm]
    table = Table(data, colWidths=col_widths, repeatRows=1)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1D4ED8')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, -1), font_name),
        ('FONTSIZE', (0, 0), (-1, -1), 9),
        ('ALIGN', (3, 1), (3, -1), 'RIGHT'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#999999')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F5F5F5')]),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 2),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 2),
    ]))
    elements.append(table)

    doc.build(elements)
    buf.seek(0)
    return buf
