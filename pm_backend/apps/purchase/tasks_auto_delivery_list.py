"""自動納入リスト送信タスク"""
import logging
import time
from datetime import date, datetime, timedelta
from io import BytesIO

from openpyxl import Workbook

logger = logging.getLogger('purchase')


def _recalculate_supplier_progress_for_auto_delivery(supplier, line, days_back, days_forward=30, product_ids=None):
    """自動納入リスト送信前に、対象仕入先ラインの進度だけを最新化する。"""
    from masters.models import BOMItem, Calendar, RoutingStep
    from orders.utils.calendar_utils import WorkingDayCalculator
    from production.inventory.inventory_calculator import recalculate_inventory_for_line
    from production.scheduler.tasks import _resolve_effective_start_date

    today = date.today()
    daiso_cal = Calendar.objects.filter(calendar_code='daiso').first()
    calc = WorkingDayCalculator(daiso_cal)
    start_date = calc.subtract_working_days(today, days_back)
    end_date = today + timedelta(days=days_forward)

    if product_ids is None:
        bom_product_ids = set(
            BOMItem.objects.filter(
                supplier_id=supplier.id,
                child_product_id__isnull=False,
            ).values_list('child_product_id', flat=True).distinct()
        )
        routing_product_ids = set(
            RoutingStep.objects.filter(
                supplier_id=supplier.id,
                output_product_id__isnull=False,
            ).values_list('output_product_id', flat=True).distinct()
        )
        product_ids = sorted(bom_product_ids | routing_product_ids)
    else:
        product_ids = sorted({int(pid) for pid in product_ids if pid})
    if not product_ids:
        return {
            'product_count': 0,
            'start_date': start_date,
            'end_date': end_date,
        }

    effective_start_date = _resolve_effective_start_date(
        line,
        start_date,
        end_date,
        today,
        include_stock_anchor=False,
        include_lt_anchor=True,
    )
    recalc_result = recalculate_inventory_for_line(
        line_id=line.id,
        start_date=start_date,
        end_date=end_date,
        include_progress=True,
        product_ids=product_ids,
        progress_only=True,
        progress_calc_start_date=effective_start_date,
    )
    recalc_result['start_date'] = start_date
    recalc_result['end_date'] = end_date
    return recalc_result


def run_auto_delivery_list_send(config_id, ignore_holiday=False):
    from notifications.models import Notification
    from masters.models import Calendar
    from orders.utils.calendar_utils import WorkingDayCalculator
    from shipping.services.email_service import EmailService

    from .models import PurchaseAutoDeliveryListConfig, SupplierOrderSchedule
    from .order_proposal_views import _generate_raw_pattern_dates
    from .views import _calc_progress_quantities, _resolve_purchase_line, _sort_delivery_list_items

    start_time = time.time()
    config = PurchaseAutoDeliveryListConfig.objects.select_related('supplier').filter(id=config_id).first()
    if not config:
        logger.error(f'自動納入リスト設定が見つかりません: config_id={config_id}')
        return

    supplier = config.supplier
    config.last_run_at = datetime.now()
    config.last_run_status = 'RUNNING'
    config.last_run_message = '実行中...'
    config.save(update_fields=['last_run_at', 'last_run_status', 'last_run_message'])

    try:
        today = date.today()
        daiso_cal = Calendar.objects.filter(calendar_code='daiso').first()
        if not daiso_cal:
            calendar_missing_message = 'ダイソウカレンダーが未設定のため、自動納品リスト実行できません'
            _finish(config, start_time, 'FAILED', calendar_missing_message)
            _notify_users(config.notify_on_failure, calendar_missing_message)
            return
        calc = WorkingDayCalculator(daiso_cal)
        if not ignore_holiday and not calc.is_working_day(today):
            holiday_message = f'自動納入リスト: 本日は休日のため、{supplier.supplier_name} 向け送信は実行しません'
            _finish(config, start_time, 'SKIPPED', holiday_message)
            _notify_users(config.notify_on_non_delivery, holiday_message)
            return
        delivery_date = calc.add_working_days(today, config.lead_time_days or 2)

        schedule = SupplierOrderSchedule.objects.filter(
            supplier_id=supplier.id, is_enabled=True,
        ).select_related('pattern').first()

        if not schedule or not schedule.pattern:
            _finish(config, start_time, 'SKIPPED', f'仕入先 {supplier.supplier_code} の納入パターンが未設定です')
            _notify_users(config.notify_on_non_delivery, f'自動納入リスト: {supplier.supplier_name} の納入パターンが未設定です')
            return

        window_start = delivery_date - timedelta(days=7)
        window_end = delivery_date + timedelta(days=90)
        pattern_dates = sorted(set(_generate_raw_pattern_dates(schedule, window_start, window_end, calc)))

        if delivery_date not in pattern_dates:
            _finish(config, start_time, 'SKIPPED', f'{delivery_date} は {supplier.supplier_code} の納入日ではありません')
            _notify_users(
                config.notify_on_non_delivery,
                f'自動納入リスト: {delivery_date} は {supplier.supplier_name} の納入日ではありません',
            )
            return

        future = [d for d in pattern_dates if d > delivery_date]
        next_delivery_date = future[0] if future else None

        if next_delivery_date:
            coverage_dates = []
            d = delivery_date
            while d < next_delivery_date:
                coverage_dates.append(d)
                d += timedelta(days=1)
        else:
            coverage_dates = [delivery_date]

        line = _resolve_purchase_line(supplier)
        if not line:
            _finish(config, start_time, 'FAILED', f'仕入先 {supplier.supplier_code} に対応するラインが見つかりません')
            _notify_users(config.notify_on_failure, f'自動納入リスト失敗: {supplier.supplier_name} に対応するラインが見つかりません')
            return

        days_back = config.progress_days_back or 7
        days_forward = config.progress_days_forward or 30
        try:
            progress_recalc_result = _recalculate_supplier_progress_for_auto_delivery(
                supplier,
                line,
                days_back,
                days_forward,
            )
            logger.info(
                '自動納入リスト送信前の進度再計算完了: supplier=%s line=%s products=%s period=%s~%s',
                supplier.supplier_code,
                line.line_code,
                progress_recalc_result.get('progress_product_count', progress_recalc_result.get('product_count', 0)),
                progress_recalc_result.get('start_date'),
                progress_recalc_result.get('end_date'),
            )
        except Exception as e:
            logger.exception(
                '自動納入リスト送信前の進度再計算エラー: supplier=%s line=%s',
                supplier.supplier_code,
                line.line_code,
            )
            _finish(config, start_time, 'FAILED', f'進度再計算エラー: {str(e)[:300]}')
            _notify_users(
                config.notify_on_failure,
                f'自動納入リスト失敗: {supplier.supplier_name} の進度再計算に失敗しました\n{str(e)[:300]}',
            )
            return

        product_map = _calc_progress_quantities(line, coverage_dates)
        items = _sort_delivery_list_items(product_map.values())

        supplier_type = getattr(supplier, 'supplier_type', 'both') or 'both'
        if supplier_type == 'outsource':
            items = [i for i in items if i['product_code'].endswith('G')]
        elif supplier_type == 'purchase':
            items = [i for i in items if not i['product_code'].endswith('G')]

        if not items:
            _finish(config, start_time, 'SUCCESS', f'{delivery_date} の納入予定品目なし（0件）')
            return

        excel_data = None
        if config.send_delivery_list_excel:
            excel_data = _generate_excel(items, delivery_date, coverage_dates, supplier)

        progress_excel = None
        progress_pdf = None
        if config.send_progress_excel:
            try:
                progress_excel = _generate_progress_excel(supplier, line, days_back, days_forward)
            except Exception as e:
                logger.warning(f'進度表Excel生成エラー（送信は続行）: {e}')
        if config.send_progress_pdf:
            try:
                progress_pdf = _generate_progress_pdf(supplier, line, days_back, days_forward)
            except Exception as e:
                logger.warning(f'進度表PDF生成エラー（送信は続行）: {e}')

        daily_items = []
        for item in items:
            daily = item.get('daily', {})
            if daily:
                for d_iso in sorted(daily.keys()):
                    d_qty = daily[d_iso]
                    if d_qty > 0:
                        daily_items.append({
                            'product_code': item['product_code'],
                            'product_name': item['product_name'],
                            'expected_qty': d_qty,
                            'delivery_date': date.fromisoformat(d_iso),
                        })
            else:
                daily_items.append({
                    'product_code': item['product_code'],
                    'product_name': item['product_name'],
                    'expected_qty': int(item['expected_qty']),
                    'delivery_date': delivery_date,
                })

        delivery_note_pdf = None
        if config.send_delivery_note_pdf:
            try:
                delivery_note_pdf = _generate_delivery_note_pdf(daily_items, delivery_date, supplier)
            except Exception as e:
                logger.warning(f'外作納品書PDF生成エラー（送信は続行）: {e}')

        to_email = (supplier.order_email or '').strip()
        if not to_email:
            _finish(config, start_time, 'FAILED', f'仕入先 {supplier.supplier_code} のメールアドレスが未設定です')
            _notify_users(config.notify_on_failure, f'自動納入リスト失敗: {supplier.supplier_name} のメールアドレスが未設定です')
            return

        cc_list = [e.strip() for e in (config.cc_emails or '').splitlines() if e.strip()]

        email_service = EmailService()
        all_attachments = []
        if excel_data:
            all_attachments.append({
                'data': excel_data,
                'filename': f'納品リスト_{supplier.supplier_code}_{delivery_date}.xlsx',
            })
        if progress_excel:
            all_attachments.append({
                'data': progress_excel,
                'filename': f'進度表_{supplier.supplier_code}_{date.today()}.xlsx',
            })
        if progress_pdf:
            all_attachments.append({
                'data': progress_pdf,
                'filename': f'進度表_{supplier.supplier_code}_{date.today()}.pdf',
            })
        if delivery_note_pdf:
            all_attachments.append({
                'data': delivery_note_pdf,
                'filename': f'外作納品書_{supplier.supplier_code}_{delivery_date}.pdf',
            })

        if not all_attachments:
            _finish(config, start_time, 'SKIPPED', f'{delivery_date} 送信ファイルがすべてOFFです')
            return

        main_attach = all_attachments[0]
        extra = all_attachments[1:] if len(all_attachments) > 1 else None

        reply_to = (config.reply_to_email or '').strip()
        reply_line = f'\n※ 返送先: {reply_to}\n（このメールは送信専用です。返信は上記アドレスへお願いいたします。）\n' if reply_to else ''
        send_timing_delivery_line = (
            f'送信タイミング: 納入日の{config.lead_time_days or 2}営業日前 '
            f'{int(config.scheduled_hour):02d}:{int(config.scheduled_minute):02d} に自動送信\n'
        )
        progress_period_line = (
            f'進度表期間: 発行日の{days_back}営業日前 ～ {days_forward}日後\n'
        )

        # 件名: 送信内容に応じて変更
        subject_parts = []
        if excel_data:
            subject_parts.append('納品リスト')
        if progress_excel or progress_pdf:
            subject_parts.append('進度表')
        if delivery_note_pdf:
            subject_parts.append('外作納品書')
        subject_label = '・'.join(subject_parts)
        if excel_data:
            subject = f'【デモ配信】【{subject_label}】納入日{delivery_date}'
        else:
            subject = f'【デモ配信】【{subject_label}】発行日{today}'

        # 本文: 送信内容に応じて構成
        body_lines = [f'{supplier.supplier_name} 御中\n', 'お世話になっております。\n']
        if excel_data:
            body_lines.append(f'納品リスト（納入日: {delivery_date}）を送付いたします。\n')
            body_lines.append('2026-07-06（月）より試運用として、自動送信を開始しております。\n')
            body_lines.append('正式運用への移行時期・運用方法は後日あらためてご相談のうえ決定いたします。それまでは、現行の発注・納入・検収方法にて運用をお願いいたします。\n')
            body_lines.append(send_timing_delivery_line)
            body_lines.append(f'対象品目: {len(items)}件')
            body_lines.append(f'カバー期間: {coverage_dates[0]} ～ {coverage_dates[-1]}\n')
            body_lines.append(progress_period_line)
            body_lines.append('添付のExcelの「確認・修正方法」シートを参照のうえ、数量確認・修正後にご返送ください。')
        else:
            body_lines.append('進度照会資料を送付いたします。\n')
            body_lines.append('2026-01-22（木）より試運用として、自動送信を開始しております。\n')
            body_lines.append('正式運用への移行時期・運用方法は後日あらためてご相談のうえ決定いたします。それまでは、現行の発注・納入・検収方法にて運用をお願いいたします。\n')
            body_lines.append(send_timing_delivery_line)
            body_lines.append(progress_period_line)
        if progress_excel or progress_pdf:
            body_lines.append('進度表を添付しておりますのでご参照ください。')
        if delivery_note_pdf:
            body_lines.append('外作納品書を添付しておりますのでご利用ください。')
        if reply_line:
            body_lines.append(reply_line)
        body_lines.append('\n------------------------------')
        body_lines.append('ダイソウ工業株式会社\n')
        body = '\n'.join(body_lines)

        result = email_service.send_email_with_attachment(
            to_emails=[to_email],
            subject=subject,
            body=body,
            attachment_data=main_attach['data'],
            attachment_filename=main_attach['filename'],
            cc_emails=cc_list if cc_list else None,
            extra_attachments=extra,
            reply_to=reply_to or None,
        )

        if result.get('success'):
            cc_info = f' CC: {", ".join(cc_list)}' if cc_list else ''
            _finish(
                config, start_time, 'SUCCESS',
                f'{delivery_date} 納入リスト送信完了 ({len(items)}件) → {to_email}{cc_info}',
            )
        else:
            _finish(config, start_time, 'FAILED', f'メール送信失敗: {result.get("message", "")}')
            _notify_users(
                config.notify_on_failure,
                f'自動納入リスト送信失敗: {supplier.supplier_name}\n{result.get("message", "")}',
            )

    except Exception as e:
        logger.exception(f'自動納入リスト送信エラー: config_id={config_id}')
        _finish(config, start_time, 'FAILED', f'エラー: {str(e)[:500]}')
        _notify_users(config.notify_on_failure, f'自動納入リスト送信エラー: {supplier.supplier_name}\n{str(e)[:300]}')


def _finish(config, start_time, status_val, message):
    duration = round(time.time() - start_time, 2)
    config.last_run_status = status_val
    config.last_run_message = message
    config.last_run_duration_seconds = duration
    config.save(update_fields=['last_run_status', 'last_run_message', 'last_run_duration_seconds'])


def _notify_users(user_m2m, message):
    from notifications.models import Notification

    user_ids = list(user_m2m.values_list('id', flat=True))
    if not user_ids:
        return
    today = date.today()
    notification = Notification.objects.create(
        title=message[:200],
        category='購買',
        domain='PURCHASE_AUTO_DELIVERY',
        valid_from=today,
        valid_to=today + timedelta(days=7),
        description=message,
        operator_name='system',
    )
    notification.target_users.set(user_ids)


def _generate_excel(items, delivery_date, coverage_dates, supplier):
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()

    # --- シート1: 納品リスト ---
    ws = wb.active
    ws.title = '納品リスト'

    DAY_NAMES = ['月', '火', '水', '木', '金', '土', '日']
    date_headers = []
    for d in coverage_dates:
        dow = DAY_NAMES[d.weekday()]
        date_headers.append(f'{d.month}/{d.day}({dow})')

    left_headers = ['品番', '品名', '移動先', '数量', '納品日']
    right_headers = ['仕入先コード', '伝票番号']
    headers = left_headers + date_headers + right_headers
    ws.append(headers)

    header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    header_font = Font(color='FFFFFF', bold=True, size=10)
    date_fill = PatternFill(start_color='D9E2F3', end_color='D9E2F3', fill_type='solid')
    date_font = Font(bold=True, size=9)
    tail_fill = PatternFill(start_color='E2EFDA', end_color='E2EFDA', fill_type='solid')
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin'),
    )

    date_start = len(left_headers) + 1
    date_end = date_start + len(date_headers) - 1
    tail_start = date_end + 1

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.border = thin_border
        cell.alignment = Alignment(horizontal='center', wrap_text=True)
        if col_idx <= len(left_headers):
            cell.fill = header_fill
            cell.font = header_font
        elif col_idx >= tail_start:
            cell.fill = tail_fill
            cell.font = Font(bold=True, size=10, color='375623')
        else:
            cell.fill = date_fill
            cell.font = date_font

    supplier_code = supplier.supplier_code or ''
    for row_idx, item in enumerate(items, 2):
        daily = item.get('daily', {})
        row_data = [
            item['product_code'],
            item['product_name'],
            item.get('transfer_destination_label', ''),
            item['expected_qty'],
            delivery_date.isoformat(),
        ]
        for d in coverage_dates:
            qty = daily.get(d.isoformat(), 0)
            row_data.append(qty if qty else '')
        row_data += [supplier_code, '']
        ws.append(row_data)
        for col_idx in range(1, len(row_data) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = thin_border
            if col_idx == 4 or (date_start <= col_idx <= date_end):
                cell.alignment = Alignment(horizontal='right')

    left_widths = [20, 30, 16, 10, 14]
    for i, w in enumerate(left_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for i in range(date_start, date_end + 1):
        ws.column_dimensions[get_column_letter(i)].width = 10
    ws.column_dimensions[get_column_letter(tail_start)].width = 16
    ws.column_dimensions[get_column_letter(tail_start + 1)].width = 14

    ws.freeze_panes = 'A2'

    # --- シート2: 確認・修正方法 ---
    ws2 = wb.create_sheet('確認・修正方法')
    instructions = [
        ['【納品リストの確認・修正方法】'],
        [''],
        ['■ 基本ルール'],
        ['', '・「納品リスト」シートの「数量」列に納入予定数量が入っています。'],
        ['', '・数量を確認し、変更がなければそのまま弊社へ返送してください。'],
        ['', '・数量を変更したい場合は「数量」列の値を直接修正してください。'],
        [''],
        ['■ 数量の修正方法'],
        ['', '1. 数量を変更する場合 → 「数量」列を修正値に書き換え'],
        ['', '2. 納入しない製品がある場合 → 「数量」列を 0 に変更'],
        ['', '3. 数量が空白の行は取込時に除外されます（取消にはなりません）'],
        [''],
        ['■ 再返送時のご注意（重要）'],
        ['', '一度返送された後に数量変更・取消が必要な場合は、修正後のこのファイルを再返送してください。'],
        ['', '・数量変更 → 「数量」列を新しい値に書き換えて再返送'],
        ['', '・取消 → 「数量」列を 0 にして再返送'],
        ['', '・行を削除したり空白にしても、前回返送した数量がそのまま残ります（取消になりません）。'],
        [''],
        ['■ 品番の追加'],
        ['', '・納品リストに品番を追加する場合は「製品リスト」シートから品番をコピーして納品リストの末行の下に追加してください。'],
        ['', '・品番を手入力すると不一致エラーの原因になります。'],
        [''],
        ['■ 納品リスト印刷用シートについて'],
        ['', '・「納品リスト印刷用」シートは納品書として印刷してご使用いただけます。'],
        ['', '・数量を変更した場合は「納品リスト印刷用」シートの数量も同様に修正してください。'],
        ['', '・貴社の自社様式の納品書をご使用いただいても構いません。'],
        [''],
        ['■ 注意事項'],
        ['', '・品番・品名・仕入先コードは変更しないでください。'],
        ['', '・このファイルをそのまま返送してください（ファイル形式を変えないこと）。'],
        ['', '・返送先はメール本文に記載のアドレスへお送りください（送信元は送信専用です）。'],
    ]
    for row in instructions:
        ws2.append(row)

    ws2.column_dimensions['A'].width = 4
    ws2.column_dimensions['B'].width = 70
    title_font = Font(bold=True, size=12)
    section_font = Font(bold=True, size=11)
    ws2.cell(row=1, column=1).font = title_font
    for r in [3, 8, 13, 19, 23, 28]:
        ws2.cell(row=r, column=1).font = section_font

    # --- シート3: 製品リスト ---
    ws3 = wb.create_sheet('製品リスト')
    ws3.append(['品番', '品名'])
    prod_header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    prod_header_font = Font(color='FFFFFF', bold=True, size=10)
    for col_idx in range(1, 3):
        cell = ws3.cell(row=1, column=col_idx)
        cell.fill = prod_header_fill
        cell.font = prod_header_font
        cell.border = thin_border

    from masters.models import Line, Product, RoutingStep
    line = Line.objects.filter(line_code=supplier.supplier_code).first()
    product_ids = set()
    if line:
        product_ids = set(
            RoutingStep.objects.filter(
                line=line,
                output_product__isnull=False,
                routing__is_active=True,
                routing__is_default=True,
            ).values_list('output_product_id', flat=True).distinct()
        )
    products = Product.objects.filter(id__in=product_ids).order_by('product_code') if product_ids else Product.objects.none()
    for row_idx, p in enumerate(products, 2):
        ws3.append([p.product_code or '', p.product_name or ''])
        for col_idx in range(1, 3):
            ws3.cell(row=row_idx, column=col_idx).border = thin_border

    ws3.column_dimensions['A'].width = 22
    ws3.column_dimensions['B'].width = 35
    ws3.freeze_panes = 'A2'

    # --- シート4: 納品リスト印刷用 ---
    from openpyxl.worksheet.properties import PageSetupProperties
    ws4 = wb.create_sheet('納品リスト印刷用')
    ws4.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws4.page_setup.orientation = 'portrait'
    ws4.page_setup.paperSize = ws4.PAPERSIZE_A4
    ws4.page_setup.fitToWidth = 1
    ws4.page_setup.fitToHeight = 0
    ws4.page_margins.left = 0.3
    ws4.page_margins.right = 0.3
    ws4.page_margins.top = 0.8
    ws4.page_margins.bottom = 0.3

    s_code = supplier.supplier_code or ''
    s_name = supplier.supplier_name or ''
    d_str = delivery_date.strftime('%Y/%m/%d')

    title_font_p = Font(bold=True, size=14)
    info_font = Font(size=10)
    info_font_bold = Font(bold=True, size=10)
    tbl_header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    tbl_header_font = Font(color='FFFFFF', bold=True, size=10)
    tbl_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin'),
    )

    ws4.merge_cells('A1:G1')
    c_title = ws4.cell(row=1, column=1, value='納 品 書')
    c_title.font = title_font_p
    c_title.alignment = Alignment(horizontal='center')

    ws4.merge_cells('A3:C3')
    ws4.cell(row=3, column=1, value='ダイソウ工業株式会社　御中').font = info_font_bold

    ws4.merge_cells('F3:G3')
    ws4.cell(row=3, column=6, value=s_name).font = info_font
    ws4.cell(row=3, column=6).alignment = Alignment(horizontal='right')

    ws4.merge_cells('F4:G4')
    ws4.cell(row=4, column=6, value=f'納品日　{d_str}').font = info_font
    ws4.cell(row=4, column=6).alignment = Alignment(horizontal='right')

    ROWS_PER_PAGE = 46
    SIGN_ROW = 57
    PAGE_BLOCK = 54
    tbl_start = 6
    p4_headers = ['No.', '品番', '品名', '数量', '単価', '金額', '備考']
    p4_widths = [5, 22, 32, 10, 10, 12, 14]

    for ci, h in enumerate(p4_headers, 1):
        cell = ws4.cell(row=tbl_start, column=ci, value=h)
        cell.font = tbl_header_font
        cell.fill = tbl_header_fill
        cell.border = tbl_border
        cell.alignment = Alignment(horizontal='center')

    for ci, w in enumerate(p4_widths, 1):
        ws4.column_dimensions[get_column_letter(ci)].width = w

    ws4.oddFooter.center.text = '&P/&Nページ'
    ws4.oddFooter.center.size = 10

    from openpyxl.worksheet.pagebreak import Break
    ws4.print_title_rows = f'{tbl_start}:{tbl_start}'

    sign_font = Font(size=9)
    sign_border_bottom = Border(bottom=Side(style='thin'))
    total_pages = max(1, -(-len(items) // ROWS_PER_PAGE))

    for page_idx in range(total_pages):
        if page_idx == 0:
            data_start = tbl_start + 1
            r_sign = SIGN_ROW
            break_row = SIGN_ROW + 1
        else:
            data_start = SIGN_ROW + 4 + (page_idx - 1) * PAGE_BLOCK
            r_sign = data_start + ROWS_PER_PAGE + 4
            break_row = r_sign + 1

        page_start = page_idx * ROWS_PER_PAGE
        page_end = min(page_start + ROWS_PER_PAGE, len(items))
        page_items = items[page_start:page_end]

        for idx, item in enumerate(page_items):
            r = data_start + idx
            ws4.cell(row=r, column=1, value=page_start + idx + 1).border = tbl_border
            ws4.cell(row=r, column=1).alignment = Alignment(horizontal='center')
            ws4.cell(row=r, column=2, value=item['product_code']).border = tbl_border
            ws4.cell(row=r, column=3, value=item['product_name']).border = tbl_border
            qty = item['expected_qty']
            c_qty = ws4.cell(row=r, column=4, value=qty)
            c_qty.border = tbl_border
            c_qty.alignment = Alignment(horizontal='right')
            ws4.cell(row=r, column=5, value='').border = tbl_border
            ws4.cell(row=r, column=5).alignment = Alignment(horizontal='right')
            ws4.cell(row=r, column=6, value='').border = tbl_border
            ws4.cell(row=r, column=6).alignment = Alignment(horizontal='right')
            ws4.cell(row=r, column=7, value='').border = tbl_border

        ws4.cell(row=r_sign, column=1, value='受領').font = sign_font
        ws4.cell(row=r_sign, column=2).border = sign_border_bottom
        ws4.cell(row=r_sign, column=5, value='納入先').font = sign_font
        ws4.cell(row=r_sign, column=5).alignment = Alignment(horizontal='right')
        ws4.merge_cells(start_row=r_sign, start_column=6, end_row=r_sign, end_column=7)
        ws4.cell(row=r_sign, column=6).border = sign_border_bottom
        ws4.cell(row=r_sign, column=7).border = sign_border_bottom
        r_date = r_sign + 2
        ws4.cell(row=r_date, column=5, value='日付').font = sign_font
        ws4.cell(row=r_date, column=5).alignment = Alignment(horizontal='right')
        ws4.merge_cells(start_row=r_date, start_column=6, end_row=r_date, end_column=7)
        ws4.cell(row=r_date, column=6).border = sign_border_bottom
        ws4.cell(row=r_date, column=7).border = sign_border_bottom

        if page_idx < total_pages - 1:
            ws4.row_breaks.append(Break(id=break_row))

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output


DAY_NAMES = ['月', '火', '水', '木', '金', '土', '日']
ROW_DEFS = [
    ('forecast', '内示'),
    ('firm', '確定'),
    ('plan', '計画'),
    ('actual', '実績'),
    ('progress', '進度'),
    ('planned_progress', '計進'),
]


def _collect_progress_data(line, days_back, days_forward=30):
    """進度表データを収集して返す。Excel/PDF共通。"""
    from collections import defaultdict
    from masters.models import Calendar
    from orders.utils.calendar_utils import WorkingDayCalculator
    from production.models import LineDemand
    from production.models_line_backlog import LineBacklog

    today = date.today()
    daiso_cal = Calendar.objects.filter(calendar_code='daiso').first()
    calc = WorkingDayCalculator(daiso_cal)
    start_date = calc.subtract_working_days(today, days_back)
    end_date = today + timedelta(days=days_forward)
    current_month_start = date(today.year, today.month, 1)
    if today.month == 12:
        next_month_start = date(today.year + 1, 1, 1)
        month_total_end = date(today.year + 1, 2, 1) - timedelta(days=1)
    else:
        next_month_start = date(today.year, today.month + 1, 1)
        if today.month + 1 == 12:
            month_total_end = date(today.year, 12, 31)
        else:
            month_total_end = date(today.year, today.month + 2, 1) - timedelta(days=1)

    date_list = []
    d = start_date
    while d <= end_date:
        date_list.append(d)
        d += timedelta(days=1)

    month_total_keys = [
        (today.year, today.month),
        (next_month_start.year, next_month_start.month),
    ]

    demand_qs = LineDemand.objects.filter(
        line=line, plan_date__gte=start_date, plan_date__lte=end_date,
    ).select_related('product')
    demand_month_total_qs = LineDemand.objects.filter(
        line=line,
        plan_date__gte=current_month_start,
        plan_date__lte=month_total_end,
    )

    demand_by_product_date = defaultdict(lambda: defaultdict(lambda: {'forecast': 0, 'firm': 0}))
    month_totals_by_product = defaultdict(lambda: {
        month_total_keys[0]: {'forecast': 0, 'firm': 0, 'plan': 0, 'actual': 0},
        month_total_keys[1]: {'forecast': 0, 'firm': 0, 'plan': 0, 'actual': 0},
    })
    product_info = {}
    for dm in demand_qs:
        pid = dm.product_id
        if pid not in product_info:
            product_info[pid] = {
                'product_code': dm.product.product_code if dm.product else dm.product_code,
                'product_name': dm.product.product_name if dm.product else '',
            }
        dd = demand_by_product_date[pid][dm.plan_date]
        dd['forecast'] += int(dm.forecast_qty or 0)
        dd['firm'] += int(dm.firm_qty or 0)
    for dm in demand_month_total_qs:
        month_key = (dm.plan_date.year, dm.plan_date.month)
        if month_key not in month_total_keys:
            continue
        pid = dm.product_id
        month_totals_by_product[pid][month_key]['forecast'] += int(dm.forecast_qty or 0)
        month_totals_by_product[pid][month_key]['firm'] += int(dm.firm_qty or 0)

    backlog_qs = LineBacklog.objects.filter(
        line=line, plan_date__gte=start_date, plan_date__lte=end_date,
    ).select_related('product')
    backlog_month_total_qs = LineBacklog.objects.filter(
        line=line,
        plan_date__gte=current_month_start,
        plan_date__lte=month_total_end,
    )

    backlog_by_product_date = defaultdict(lambda: defaultdict(lambda: {
        'plan': 0, 'actual': 0, 'progress': 0, 'planned_progress': 0,
    }))
    for bl in backlog_qs:
        pid = bl.product_id
        if pid not in product_info and bl.product:
            product_info[pid] = {
                'product_code': bl.product.product_code,
                'product_name': bl.product.product_name,
            }
        bd = backlog_by_product_date[pid][bl.plan_date]
        bd['plan'] += int(bl.plan_qty or 0)
        bd['actual'] += int(bl.actual_qty or 0)
        bd['progress'] += int(bl.progress_qty or 0)
        bd['planned_progress'] += int(getattr(bl, 'planned_progress_qty', 0) or 0)
    for bl in backlog_month_total_qs:
        month_key = (bl.plan_date.year, bl.plan_date.month)
        if month_key not in month_total_keys:
            continue
        pid = bl.product_id
        month_totals_by_product[pid][month_key]['plan'] += int(bl.plan_qty or 0)
        month_totals_by_product[pid][month_key]['actual'] += int(bl.actual_qty or 0)

    if not product_info:
        return None

    # 繰越: 開始日前日の進度・計進を取得
    prev_date = start_date - timedelta(days=1)
    carryover_qs = LineBacklog.objects.filter(
        line=line, plan_date=prev_date, product_id__in=product_info.keys(),
    )
    carryover_by_pid = {}
    for bl in carryover_qs:
        pid = bl.product_id
        if pid not in carryover_by_pid:
            carryover_by_pid[pid] = {'progress': 0, 'planned_progress': 0}
        carryover_by_pid[pid]['progress'] += int(bl.progress_qty or 0)
        carryover_by_pid[pid]['planned_progress'] += int(getattr(bl, 'planned_progress_qty', 0) or 0)

    sorted_pids = sorted(product_info.keys(), key=lambda pid: product_info[pid]['product_code'])
    return {
        'today': today,
        'start_date': start_date,
        'end_date': end_date,
        'date_list': date_list,
        'product_info': product_info,
        'sorted_pids': sorted_pids,
        'demand_by_product_date': demand_by_product_date,
        'backlog_by_product_date': backlog_by_product_date,
        'carryover_by_pid': carryover_by_pid,
        'month_total_keys': month_total_keys,
        'month_totals_by_product': month_totals_by_product,
    }

def _generate_progress_excel(supplier, line, days_back, days_forward=30):
    """仕入先の全製品の進度表Excelを生成"""
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    data = _collect_progress_data(line, days_back, days_forward)
    if not data:
        return None

    today = data['today']
    start_date = data['start_date']
    end_date = data['end_date']
    date_list = data['date_list']
    product_info = data['product_info']
    sorted_pids = data['sorted_pids']
    demand_by_product_date = data['demand_by_product_date']
    backlog_by_product_date = data['backlog_by_product_date']
    carryover_by_pid = data['carryover_by_pid']
    month_keys = data['month_total_keys']
    month_totals_by_product = data['month_totals_by_product']

    wb = Workbook()
    ws = wb.active
    ws.title = '進度表'

    FONT_SIZE = 9
    base_font = Font(size=FONT_SIZE)
    header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    header_font = Font(color='FFFFFF', bold=True, size=FONT_SIZE)
    label_fill = PatternFill(start_color='F2F2F2', end_color='F2F2F2', fill_type='solid')
    label_font = Font(bold=True, size=FONT_SIZE)
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin'),
    )
    today_fill = PatternFill(start_color='D6EAFF', end_color='D6EAFF', fill_type='solid')
    negative_font = Font(color='FF0000', size=FONT_SIZE)
    positive_font = Font(color='0000FF', size=FONT_SIZE)
    weekend_fill = PatternFill(start_color='FFF0F0', end_color='FFF0F0', fill_type='solid')

    current_row = 1

    ws.cell(row=current_row, column=1, value=f'進度表 — {supplier.supplier_code} {supplier.supplier_name}')
    ws.cell(row=current_row, column=1).font = Font(bold=True, size=12)
    current_row += 1
    ws.cell(row=current_row, column=1, value=f'発行日: {today}　期間: {start_date} ～ {end_date}')
    ws.cell(row=current_row, column=1).font = Font(size=9, color='666666')
    current_row += 1
    legend_cell = ws.cell(row=current_row, column=1, value='■ 当日')
    legend_cell.font = Font(size=8, color='4A90D9')
    ws.cell(row=current_row, column=1).fill = PatternFill(start_color='D6EAFF', end_color='D6EAFF', fill_type='solid')
    ws.cell(row=current_row, column=2, value='■ 休日')
    ws.cell(row=current_row, column=2).font = Font(size=8, color='CC6666')
    ws.cell(row=current_row, column=2).fill = PatternFill(start_color='FFD9D9', end_color='FFD9D9', fill_type='solid')
    current_row += 1

    row_explanations = [
        '内示: 顧客からの予定需要です。まだ確定していない見込み数量です。',
        '確定: 顧客から正式確定された需要です。出荷・生産として確実に必要な数量です。',
        '計画: その日納入予定の数量です。',
        '実績: その日納入された数量です。',
        '進度: 納入が需要に対してどれだけ先行/遅れしているかを表す差分です。',
        '計進: 納入計画ベースでの進度です。納入計画数量まで含めた進み具合を見るための数です。（注: 発行日の納入実績は計算されません）',
    ]
    for text in row_explanations:
        ws.cell(row=current_row, column=1, value=text).font = Font(size=8, color='666666')
        current_row += 1

    current_row += 1

    for pid in sorted_pids:
        info = product_info[pid]
        demands = demand_by_product_date[pid]
        backlogs = backlog_by_product_date[pid]

        ws.cell(row=current_row, column=1, value=info['product_code'])
        ws.cell(row=current_row, column=1).font = Font(bold=True, size=10)
        current_row += 1
        ws.cell(row=current_row, column=1, value=info['product_name'])
        ws.cell(row=current_row, column=1).font = Font(size=9, color='555555')

        date_header_row = current_row
        sub_header_row = current_row + 1
        month_col_start = 3
        for offset, (_, month) in enumerate(month_keys):
            header_cell = ws.cell(row=date_header_row, column=month_col_start + offset, value=f'{month}月')
            header_cell.font = header_font
            header_cell.fill = header_fill
            header_cell.alignment = Alignment(horizontal='center')
            header_cell.border = thin_border

            sub_cell = ws.cell(row=sub_header_row, column=month_col_start + offset, value='合計数')
            sub_cell.font = header_font
            sub_cell.fill = header_fill
            sub_cell.alignment = Alignment(horizontal='center')
            sub_cell.border = thin_border

        carry_col = month_col_start + len(month_keys)
        carry_cell = ws.cell(row=date_header_row, column=carry_col, value='繰越')
        carry_cell.font = header_font
        carry_cell.fill = header_fill
        carry_cell.alignment = Alignment(horizontal='center')
        carry_cell.border = thin_border
        ws.merge_cells(
            start_row=date_header_row,
            start_column=carry_col,
            end_row=sub_header_row,
            end_column=carry_col,
        )
        DATE_COL_START = carry_col + 1
        for col_idx, dt in enumerate(date_list, DATE_COL_START):
            dow = DAY_NAMES[dt.weekday()]
            header_cell = ws.cell(row=date_header_row, column=col_idx, value=f'{dt.month}/{dt.day}')
            header_cell.font = header_font
            header_cell.fill = header_fill
            header_cell.alignment = Alignment(horizontal='center')
            header_cell.border = thin_border

            sub_cell = ws.cell(row=sub_header_row, column=col_idx, value=dow)
            sub_cell.font = header_font
            sub_cell.fill = header_fill
            sub_cell.alignment = Alignment(horizontal='center')
            sub_cell.border = thin_border
        current_row += 2

        carryover = carryover_by_pid.get(pid, {})
        for row_key, row_label in ROW_DEFS:
            ws.cell(row=current_row, column=1, value='')
            label_cell = ws.cell(row=current_row, column=2, value=row_label)
            label_cell.fill = label_fill
            label_cell.font = label_font
            label_cell.border = thin_border

            for offset, month_key in enumerate(month_keys):
                month_total = None
                if row_key not in ('progress', 'planned_progress'):
                    month_total = month_totals_by_product.get(pid, {}).get(month_key, {}).get(row_key, 0)

                month_cell = ws.cell(
                    row=current_row,
                    column=month_col_start + offset,
                    value=month_total if month_total not in (None, 0) else '',
                )
                month_cell.border = thin_border
                month_cell.alignment = Alignment(horizontal='right')
                month_cell.font = base_font

            # 繰越列
            cv = carryover.get(row_key, 0) if row_key in ('progress', 'planned_progress') else ''
            cc = ws.cell(row=current_row, column=carry_col, value=cv if cv else '')
            cc.border = thin_border
            cc.alignment = Alignment(horizontal='right')
            cc.font = base_font
            if row_key in ('progress', 'planned_progress') and cv:
                cc.font = negative_font if cv < 0 else (positive_font if cv > 0 else base_font)

            for col_idx, dt in enumerate(date_list, DATE_COL_START):
                if row_key in ('forecast', 'firm'):
                    val = demands[dt].get(row_key, 0) if dt in demands else 0
                else:
                    val = backlogs[dt].get(row_key, 0) if dt in backlogs else 0

                cell = ws.cell(row=current_row, column=col_idx, value=val if val else '')
                cell.border = thin_border
                cell.alignment = Alignment(horizontal='right')
                cell.font = base_font

                is_weekend = dt.weekday() >= 5
                is_today = dt == today

                if row_key in ('progress', 'planned_progress') and val:
                    if val < 0:
                        cell.font = negative_font
                    elif val > 0:
                        cell.font = positive_font

                if is_today:
                    cell.fill = today_fill
                elif is_weekend:
                    cell.fill = weekend_fill

            current_row += 1

        current_row += 1

    ws.column_dimensions['A'].width = 18
    ws.column_dimensions['B'].width = 10
    for col_idx in range(month_col_start, carry_col):
        ws.column_dimensions[get_column_letter(col_idx)].width = 8
    ws.column_dimensions['C'].width = 8
    for col_idx in range(DATE_COL_START, len(date_list) + DATE_COL_START):
        ws.column_dimensions[get_column_letter(col_idx)].width = 7
    ws.column_dimensions[get_column_letter(carry_col)].width = 8
    ws.freeze_panes = f'{get_column_letter(DATE_COL_START)}11'

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def _generate_progress_pdf(supplier, line, days_back, days_forward=30):
    """仕入先の全製品の進度表PDFを生成（横向き）"""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.units import mm
    from reportlab.pdfgen import canvas as pdf_canvas
    from shipping.services.shipping_pdf_generator import register_japanese_fonts

    data = _collect_progress_data(line, days_back, days_forward)
    if not data:
        return None

    register_japanese_fonts()

    today = data['today']
    start_date = data['start_date']
    end_date = data['end_date']
    date_list = data['date_list']
    product_info = data['product_info']
    sorted_pids = data['sorted_pids']
    demand_by_product_date = data['demand_by_product_date']
    backlog_by_product_date = data['backlog_by_product_date']
    carryover_by_pid = data['carryover_by_pid']
    month_keys = data['month_total_keys']
    month_totals_by_product = data['month_totals_by_product']

    page_w, page_h = landscape(A4)
    margin_left = 5 * mm
    margin_right = 3 * mm
    margin_top = 12 * mm
    margin_bottom = 6 * mm

    CARRY_W = 9 * mm
    MONTH_W = 8 * mm
    LABEL_W = 8 * mm
    LEFT_W = (MONTH_W * len(month_keys)) + LABEL_W + CARRY_W
    DATE_W = 7.5 * mm
    ROW_H = 3.2 * mm
    HEADER_H = 4.5 * mm
    FONT_NAME = 'MSGothic'
    FONT_SIZE = 5.5
    HEADER_FONT_SIZE = 5

    available_w = page_w - margin_left - margin_right - LEFT_W
    DATES_PER_PAGE = int(available_w / DATE_W)

    buf = BytesIO()
    c = pdf_canvas.Canvas(buf, pagesize=landscape(A4))

    date_chunks = [date_list[i:i + DATES_PER_PAGE] for i in range(0, len(date_list), DATES_PER_PAGE)]
    product_block_h = (len(ROW_DEFS) + 2) * ROW_H + HEADER_H
    usable_h = page_h - margin_top - margin_bottom - 10 * mm
    products_per_page = max(1, int(usable_h / product_block_h))

    total_pages = 0
    for _ in date_chunks:
        total_pages += max(1, -(-len(sorted_pids) // products_per_page))

    page_num = 0
    for date_chunk in date_chunks:
        pid_chunks = [sorted_pids[i:i + products_per_page] for i in range(0, len(sorted_pids), products_per_page)]

        for pid_chunk in pid_chunks:
            page_num += 1

            c.setFont(FONT_NAME, 9)
            c.drawString(margin_left, page_h - 7 * mm, f'進度表 — {supplier.supplier_code} {supplier.supplier_name}')
            c.setFont(FONT_NAME, 5)
            c.setFillColor(colors.grey)
            c.drawString(margin_left, page_h - 10.5 * mm, f'発行日: {today}　期間: {start_date} ～ {end_date}')
            c.drawRightString(page_w - margin_right, page_h - 7 * mm, f'PAGE {page_num}/{total_pages}')

            legend_y = page_h - 10.5 * mm
            legend_x = page_w - margin_right - 80 * mm
            box = 2.5 * mm
            c.setFillColor(colors.Color(0.84, 0.92, 1.0))
            c.rect(legend_x, legend_y - 0.5 * mm, box, box, fill=1, stroke=1)
            c.setFillColor(colors.grey)
            c.drawString(legend_x + box + 1 * mm, legend_y, '当日')
            legend_x += 16 * mm
            c.setFillColor(colors.Color(1, 0.85, 0.85))
            c.rect(legend_x, legend_y - 0.5 * mm, box, box, fill=1, stroke=1)
            c.setFillColor(colors.grey)
            c.drawString(legend_x + box + 1 * mm, legend_y, '休日')
            c.setFillColor(colors.black)

            y = page_h - margin_top - 2 * mm

            for pid in pid_chunk:
                info = product_info[pid]
                demands = demand_by_product_date[pid]
                backlogs = backlog_by_product_date[pid]

                # 品番行
                x = margin_left
                c.setFillColor(colors.black)
                c.setFont(FONT_NAME, 6)
                c.drawString(x, y - ROW_H * 0.75, info['product_code'])
                y -= ROW_H
                # 品名行
                c.setFont(FONT_NAME, 5)
                c.setFillColor(colors.Color(0.3, 0.3, 0.3))
                c.drawString(x + 1 * mm, y - ROW_H * 0.75, info['product_name'] or '')
                c.setFillColor(colors.black)
                y -= ROW_H

                GRID_COLOR = colors.Color(0.5, 0.5, 0.5)
                c.setLineWidth(0.25)
                WEEKEND_BG = colors.Color(1, 0.85, 0.85)
                TODAY_BG = colors.Color(0.84, 0.92, 1.0)

                month_x = margin_left
                label_x = month_x + (MONTH_W * len(month_keys))

                # 繰越ヘッダー
                for idx, (_, month) in enumerate(month_keys):
                    col_x = month_x + (MONTH_W * idx)
                    c.setStrokeColor(GRID_COLOR)
                    c.rect(col_x, y - HEADER_H, MONTH_W, HEADER_H, fill=0, stroke=1)
                    c.setFillColor(colors.black)
                    c.setFont(FONT_NAME, HEADER_FONT_SIZE)
                    c.drawCentredString(col_x + MONTH_W / 2, y - 3 * mm, f'{month}月')

                carry_x = label_x + LABEL_W
                c.setStrokeColor(GRID_COLOR)
                c.rect(carry_x, y - HEADER_H, CARRY_W, HEADER_H, fill=0, stroke=1)
                c.setFillColor(colors.black)
                c.setFont(FONT_NAME, HEADER_FONT_SIZE)
                c.drawCentredString(carry_x + CARRY_W / 2, y - 3 * mm, '繰越')

                # 日付ヘッダー
                x = margin_left + LEFT_W
                prev_month = None

                for dt in date_chunk:
                    is_today = dt == today
                    is_weekend = dt.weekday() >= 5

                    if is_weekend:
                        c.setFillColor(WEEKEND_BG)
                        c.rect(x, y - HEADER_H, DATE_W, HEADER_H, fill=1, stroke=0)
                    if is_today:
                        c.setFillColor(TODAY_BG)
                        c.rect(x, y - HEADER_H, DATE_W, HEADER_H, fill=1, stroke=0)

                    c.setStrokeColor(GRID_COLOR)
                    c.rect(x, y - HEADER_H, DATE_W, HEADER_H, fill=0, stroke=1)

                    c.setFillColor(colors.black)
                    c.setFont(FONT_NAME, HEADER_FONT_SIZE)
                    label = f'{dt.day}日'
                    if prev_month is None or dt.month != prev_month:
                        label = f'{dt.month}/{dt.day}'
                    prev_month = dt.month
                    dow = DAY_NAMES[dt.weekday()]
                    c.drawCentredString(x + DATE_W / 2, y - 2.2 * mm, label)
                    c.setFont(FONT_NAME, 4)
                    c.drawCentredString(x + DATE_W / 2, y - HEADER_H + 0.5 * mm, dow)
                    x += DATE_W

                y -= HEADER_H

                # データ行
                carryover = carryover_by_pid.get(pid, {})
                for row_key, row_label in ROW_DEFS:
                    # ラベル
                    lx = label_x
                    c.setFillColor(colors.Color(0.9, 0.9, 0.9))
                    c.rect(lx, y - ROW_H, LABEL_W, ROW_H, fill=1, stroke=0)
                    c.setStrokeColor(GRID_COLOR)
                    c.rect(lx, y - ROW_H, LABEL_W, ROW_H, fill=0, stroke=1)
                    c.setFillColor(colors.black)
                    c.setFont(FONT_NAME, FONT_SIZE)
                    c.drawString(lx + 0.5 * mm, y - ROW_H * 0.78, row_label)

                    # 当月/翌月合計セル
                    for idx, month_key in enumerate(month_keys):
                        col_x = month_x + (MONTH_W * idx)
                        c.setStrokeColor(GRID_COLOR)
                        c.rect(col_x, y - ROW_H, MONTH_W, ROW_H, fill=0, stroke=1)
                        if row_key not in ('progress', 'planned_progress'):
                            month_total = month_totals_by_product.get(pid, {}).get(month_key, {}).get(row_key, 0)
                            if month_total:
                                c.setFillColor(colors.black)
                                c.setFont(FONT_NAME, FONT_SIZE)
                                c.drawRightString(col_x + MONTH_W - 0.5 * mm, y - ROW_H * 0.78, str(month_total))

                    # 繰越セル
                    c.setStrokeColor(GRID_COLOR)
                    c.rect(carry_x, y - ROW_H, CARRY_W, ROW_H, fill=0, stroke=1)
                    cv = carryover.get(row_key, 0) if row_key in ('progress', 'planned_progress') else 0
                    if cv:
                        if cv < 0:
                            c.setFillColor(colors.red)
                        elif cv > 0:
                            c.setFillColor(colors.blue)
                        else:
                            c.setFillColor(colors.black)
                        c.setFont(FONT_NAME, FONT_SIZE)
                        c.drawRightString(carry_x + CARRY_W - 0.5 * mm, y - ROW_H * 0.78, str(cv))
                        c.setFillColor(colors.black)

                    # 日付セル
                    x = margin_left + LEFT_W
                    for dt in date_chunk:
                        is_today = dt == today
                        is_weekend = dt.weekday() >= 5

                        if is_today:
                            c.setFillColor(TODAY_BG)
                            c.rect(x, y - ROW_H, DATE_W, ROW_H, fill=1, stroke=0)
                        elif is_weekend:
                            c.setFillColor(WEEKEND_BG)
                            c.rect(x, y - ROW_H, DATE_W, ROW_H, fill=1, stroke=0)

                        c.setStrokeColor(GRID_COLOR)
                        c.rect(x, y - ROW_H, DATE_W, ROW_H, fill=0, stroke=1)

                        if row_key in ('forecast', 'firm'):
                            val = demands[dt].get(row_key, 0) if dt in demands else 0
                        else:
                            val = backlogs[dt].get(row_key, 0) if dt in backlogs else 0

                        if val:
                            if row_key in ('progress', 'planned_progress') and val < 0:
                                c.setFillColor(colors.red)
                            elif row_key in ('progress', 'planned_progress') and val > 0:
                                c.setFillColor(colors.blue)
                            else:
                                c.setFillColor(colors.black)
                            c.setFont(FONT_NAME, FONT_SIZE)
                            c.drawRightString(x + DATE_W - 0.5 * mm, y - ROW_H * 0.78, str(val))
                            c.setFillColor(colors.black)

                        x += DATE_W
                    y -= ROW_H

                # 製品間の区切り線
                table_right = margin_left + LEFT_W + len(date_chunk) * DATE_W
                c.setLineWidth(0.5)
                c.setStrokeColor(colors.black)
                c.line(margin_left, y, table_right, y)
                c.setLineWidth(0.25)
                y -= 1.5 * mm

            c.showPage()

    c.save()
    buf.seek(0)
    return buf


def _generate_delivery_note_pdf(items, delivery_date, supplier):
    """外作納品書PDFを生成（品目ごとに納品書・受領書・購入先控の3枚組、1ページ3品目）"""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.units import mm
    from reportlab.pdfgen import canvas as pdf_canvas
    from shipping.services.shipping_pdf_generator import register_japanese_fonts

    register_japanese_fonts()
    FONT = 'MSGothic'

    page_w, page_h = landscape(A4)
    ML, MR, MT, MB = 10 * mm, 10 * mm, 3 * mm, 3 * mm
    usable_w = page_w - ML - MR
    usable_h = page_h - MT - MB
    PER_PAGE = 3
    BLOCK_H = usable_h / PER_PAGE
    SEP = 5 * mm
    ITEM_H = BLOCK_H - SEP

    GAP = 4 * mm
    NOHIN_W = usable_w * 0.50
    SIDE_W = (usable_w - NOHIN_W - GAP * 2) / 2

    s_code = supplier.supplier_code or ''
    s_name = supplier.supplier_name or ''

    buf = BytesIO()
    c = pdf_canvas.Canvas(buf, pagesize=landscape(A4))
    total_pages = max(1, -(-len(items) // PER_PAGE))

    for pi in range(total_pages):
        page_items = items[pi * PER_PAGE:(pi + 1) * PER_PAGE]
        for ii, item in enumerate(page_items):
            yt = page_h - MT - ii * BLOCK_H
            pc = item['product_code']
            pn = item['product_name']
            qty = int(item['expected_qty'])
            item_date = item.get('delivery_date', delivery_date)
            d_ymd = item_date.strftime('%Y/%m/%d')
            d_mmdd = f'{item_date.month:02d}/{item_date.day:02d}'
            qr_data = f'{pc},{item_date.isoformat()},{qty}'

            _dn_nohin(c, ML, yt, NOHIN_W, ITEM_H,
                      pc, pn, qty, d_ymd, d_mmdd, s_code, s_name, qr_data, FONT)
            _dn_side(c, ML + NOHIN_W + GAP, yt, SIDE_W, ITEM_H,
                     pc, pn, qty, d_ymd, d_mmdd, s_code, s_name, qr_data, FONT, '受領書')
            _dn_side(c, ML + NOHIN_W + GAP + SIDE_W + GAP, yt, SIDE_W, ITEM_H,
                     pc, pn, qty, d_ymd, d_mmdd, s_code, s_name, qr_data, FONT, '購入先控')

            if ii < len(page_items) - 1:
                sy = yt - BLOCK_H + SEP / 2
                c.saveState()
                c.setDash(4, 3)
                c.setLineWidth(0.3)
                c.setStrokeColor(colors.HexColor('#888888'))
                c.line(ML, sy, page_w - MR, sy)
                c.restoreState()

        c.showPage()

    c.save()
    buf.seek(0)
    return buf


def _dn_nohin(c, x0, yt, W, H, pc, pn, qty, d_ymd, d_mmdd, s_code, s_name, qr_data, F):
    """外作納品書 — 納品書セクション描画"""
    from reportlab.lib.units import mm
    from reportlab.graphics.barcode.qr import QrCodeWidget
    from reportlab.graphics import renderPDF
    from reportlab.graphics.shapes import Drawing

    c.saveState()
    c.setLineWidth(0.4)

    TH = 6 * mm
    R0, R1, R2 = 10 * mm, 14 * mm, 10 * mm
    GH = H - TH
    R3 = GH - R0 - R1 - R2

    gy = yt - TH
    y0 = gy
    y1 = gy - R0
    y2 = y1 - R1
    y3 = y2 - R2
    yb = y3 - R3

    C1 = 63 * mm
    xr = x0 + C1
    RW = W - C1

    # --- Title ---
    c.setFont(F, 10)
    c.drawString(x0 + 1 * mm, yt - 4.5 * mm, '納　品　書')
    c.setFont(F, 10)
    c.drawString(x0 + 35 * mm, yt - 4.5 * mm, 'ダイソウ工業株式会社　御中')
    c.setFont(F, 6.5)
    c.drawRightString(x0 + W - 1 * mm, yt - 4.5 * mm, f'発行日 {d_ymd}')

    # --- Outer border ---
    c.rect(x0, yb, W, GH)

    # --- Row 0: 取引先 / 搬入場所 / 注文No. ---
    bw = RW * 0.38
    c.line(x0, y1, x0 + C1 + bw, y1)
    c.line(xr, y0, xr, y2)
    c.line(xr + bw, y0, xr + bw, y2)

    c.setFont(F, 5)
    c.drawString(x0 + 1 * mm, y0 - 3 * mm, '取引先')
    c.setFont(F, 7)
    s_label = f'{s_name}'
    if c.stringWidth(s_label, F, 7) > C1 - 14 * mm:
        c.setFont(F, 5.5)
    c.drawString(x0 + 12 * mm, y0 - 3 * mm, s_label)
    c.setFont(F, 7)
    c.drawString(x0 + 1 * mm, y0 - 8 * mm, s_code)

    c.setFont(F, 5)
    c.drawString(xr + 1 * mm, y0 - 3 * mm, '搬入場所')
    c.drawString(xr + bw + 1 * mm, y0 - 3 * mm, '注文No.')

    # --- Row 1: 部品番号 / 搬入月日 / 搬入数 ---
    c.line(x0, y2, x0 + W, y2)
    hw = bw / 2
    c.line(xr + hw, y1, xr + hw, y2)

    c.setFont(F, 5)
    c.drawString(x0 + 1 * mm, y1 - 3 * mm, '部 品 番 号')
    pc_sz = 9
    if c.stringWidth(pc, F, pc_sz) > C1 - 3 * mm:
        pc_sz = 7
    c.setFont(F, pc_sz)
    c.drawString(x0 + 1 * mm, y1 - 10 * mm, pc)

    c.setFont(F, 5)
    c.drawString(xr + 1 * mm, y1 - 3 * mm, '搬入月日')
    c.setFont(F, 12)
    c.drawCentredString(xr + hw / 2, y1 - 11 * mm, d_mmdd)

    c.setFont(F, 5)
    c.drawString(xr + hw + 1 * mm, y1 - 3 * mm, '搬入数')
    c.setFont(F, 12)
    c.drawCentredString(xr + hw + hw / 2, y1 - 11 * mm, str(qty))

    # --- Row 2: 部品名称 / 日付変更時 / 数量変更時 / 廃棄数 ---
    c.line(x0, y3, x0 + W, y3)
    tw = RW / 3
    c.line(xr, y2, xr, y3)
    c.line(xr + tw, y2, xr + tw, y3)
    c.line(xr + 2 * tw, y2, xr + 2 * tw, y3)

    c.setFont(F, 5)
    c.drawString(x0 + 1 * mm, y2 - 3 * mm, '部 品 名 称')
    pn_sz = 8
    if c.stringWidth(pn, F, pn_sz) > C1 - 3 * mm:
        pn_sz = 6
    c.setFont(F, pn_sz)
    c.drawString(x0 + 1 * mm, y2 - 8.5 * mm, pn)

    c.setFont(F, 5)
    c.drawString(xr + 1 * mm, y2 - 3 * mm, '日付変更時')
    c.drawString(xr + tw + 1 * mm, y2 - 3 * mm, '数量変更時')
    c.drawString(xr + 2 * tw + 1 * mm, y2 - 3 * mm, '廃棄数')

    # --- Row 3: 得意先 / QR / 荷姿+発行受領入力 ---
    qrw = 20 * mm
    got_w = C1 - qrw

    c.line(x0 + got_w, y3, x0 + got_w, yb)
    c.line(xr, y3, xr, yb)

    c.setFont(F, 5)
    c.drawString(x0 + 0.5 * mm, y3 - 3 * mm, '得意先')

    # QR code（欄いっぱいに大きく表示）
    qr_sz = min(R3 - 2 * mm, qrw - 2 * mm)
    if qr_sz > 4 * mm:
        try:
            qr = QrCodeWidget(qr_data, barWidth=qr_sz, barHeight=qr_sz)
            d = Drawing(qr_sz, qr_sz)
            d.add(qr)
            qr_x = x0 + got_w + (qrw - qr_sz) / 2
            qr_y = yb + (R3 - qr_sz) / 2
            renderPDF.draw(d, c, qr_x, qr_y)
        except Exception:
            pass

    sw = 8 * mm
    nw = RW - sw * 3
    c.line(xr + nw, y3, xr + nw, yb)
    c.line(xr + nw + sw, y3, xr + nw + sw, yb)
    c.line(xr + nw + sw * 2, y3, xr + nw + sw * 2, yb)

    c.setFont(F, 5)
    c.drawString(xr + 0.5 * mm, y3 - 3 * mm, '荷姿(入り数×台数)')
    c.setFont(F, 6)
    lh = 4 * mm
    c.drawString(xr + 0.5 * mm, y3 - 8 * mm, '１ポリ(　　×　　)')
    c.drawString(xr + 0.5 * mm, y3 - 8 * mm - lh, '２アミ(　　×　　)')
    c.drawString(xr + 0.5 * mm, y3 - 8 * mm - lh * 2, '３専用(　　×　　)')

    c.setFont(F, 5)
    for i, lab in enumerate(['発行', '受領', '入力']):
        sx = xr + nw + sw * i
        c.drawCentredString(sx + sw / 2, y3 - 3 * mm, lab)

    c.restoreState()


def _dn_side(c, x0, yt, W, H, pc, pn, qty, d_ymd, d_mmdd, s_code, s_name, qr_data, F, title):
    """外作納品書 — 受領書/購入先控セクション描画"""
    from reportlab.lib.units import mm

    c.saveState()
    c.setLineWidth(0.4)

    TH = 6 * mm
    R0, R1, R2, R3, R4 = 8 * mm, 12 * mm, 7 * mm, 9 * mm, 7 * mm
    GH = H - TH
    R5 = GH - R0 - R1 - R2 - R3 - R4

    gy = yt - TH
    y0 = gy
    y1 = y0 - R0
    y2 = y1 - R1
    y3 = y2 - R2
    y4 = y3 - R3
    y5 = y4 - R4
    yb = y5 - R5

    # --- Title ---
    c.setFont(F, 7)
    c.drawString(x0 + 0.5 * mm, yt - 4 * mm, title)
    c.setFont(F, 5.5)
    t_offset = 22 * mm if title == '購入先控' else 15 * mm
    c.drawString(x0 + t_offset, yt - 4 * mm, 'ダイソウ工業㈱')
    c.drawRightString(x0 + W - 0.5 * mm, yt - 4 * mm, d_ymd)

    # --- Outer border ---
    c.rect(x0, yb, W, GH)

    # --- Row 0: 取引先 / 注文No. ---
    hw0 = W * 0.65
    c.line(x0, y1, x0 + W, y1)
    c.line(x0 + hw0, y0, x0 + hw0, y1)
    c.setFont(F, 5)
    c.drawString(x0 + 0.5 * mm, y0 - 3 * mm, '取引先')
    c.drawString(x0 + hw0 + 0.5 * mm, y0 - 3 * mm, '注文No.')
    s_lbl = f'{s_code}  {s_name}'
    s_sz = 5.5
    if c.stringWidth(s_lbl, F, s_sz) > hw0 - 2 * mm:
        s_sz = 4
    c.setFont(F, s_sz)
    c.drawString(x0 + 0.5 * mm, y0 - 6.5 * mm, s_lbl)

    # --- Row 1: 部品番号 ---
    c.line(x0, y2, x0 + W, y2)
    c.setFont(F, 5)
    c.drawString(x0 + 0.5 * mm, y1 - 3 * mm, '部 品 番 号')
    pc_sz = 7
    if c.stringWidth(pc, F, pc_sz) > W - 18 * mm:
        pc_sz = 5
    c.setFont(F, pc_sz)
    c.drawString(x0 + 16 * mm, y1 - 3 * mm, pc)

    # --- Row 2: 部品名称 ---
    c.line(x0, y3, x0 + W, y3)
    c.setFont(F, 5)
    c.drawString(x0 + 0.5 * mm, y2 - 3 * mm, '部 品 名 称')
    pn_sz = 6
    if c.stringWidth(pn, F, pn_sz) > W - 2 * mm:
        pn_sz = 4.5
    c.setFont(F, pn_sz)
    c.drawString(x0 + 10 * mm, y2 - 3 * mm - 3 * mm, pn)

    # --- Row 3: 搬入月日 / 搬入数 ---
    c.line(x0, y4, x0 + W, y4)
    mid = W * 0.5
    c.line(x0 + mid, y3, x0 + mid, y4)
    c.setFont(F, 5)
    c.drawString(x0 + 0.5 * mm, y3 - 3 * mm, '搬入月日')
    c.drawString(x0 + mid + 0.5 * mm, y3 - 3 * mm, '搬入数')
    c.setFont(F, 10)
    c.drawCentredString(x0 + mid / 2, y3 - 7.5 * mm, d_mmdd)
    c.drawCentredString(x0 + mid + (W - mid) / 2, y3 - 7.5 * mm, str(qty))

    # --- Row 4: 日付変更時 / 数量変更時 / 廃棄数 ---
    c.line(x0, y5, x0 + W, y5)
    tw = W / 3
    c.line(x0 + tw, y4, x0 + tw, y5)
    c.line(x0 + 2 * tw, y4, x0 + 2 * tw, y5)
    c.setFont(F, 5)
    c.drawString(x0 + 0.5 * mm, y4 - 3 * mm, '日付変更時')
    c.drawString(x0 + tw + 0.5 * mm, y4 - 3 * mm, '数量変更時')
    c.drawString(x0 + 2 * tw + 0.5 * mm, y4 - 3 * mm, '廃棄数')

    # --- Row 5: QR / 当社 / 取引先 ---
    from reportlab.graphics.barcode.qr import QrCodeWidget
    from reportlab.graphics.shapes import Drawing
    from reportlab.graphics import renderPDF as _renderPDF
    icw = W * 0.45
    rest = W - icw
    c.line(x0 + icw, y5, x0 + icw, yb)
    c.line(x0 + icw + rest / 2, y5, x0 + icw + rest / 2, yb)

    qr_sz = min(R5 - 2 * mm, icw - 2 * mm)
    if qr_sz > 4 * mm:
        try:
            qr = QrCodeWidget(qr_data, barWidth=qr_sz, barHeight=qr_sz)
            d = Drawing(qr_sz, qr_sz)
            d.add(qr)
            qr_x = x0 + (icw - qr_sz) / 2
            qr_y = yb + (R5 - qr_sz) / 2
            _renderPDF.draw(d, c, qr_x, qr_y)
        except Exception:
            pass

    c.setFont(F, 5)
    c.drawCentredString(x0 + icw + rest / 4, y5 - 3 * mm, '当社')
    c.drawCentredString(x0 + icw + 3 * rest / 4, y5 - 3 * mm, '取引先')

    c.restoreState()
