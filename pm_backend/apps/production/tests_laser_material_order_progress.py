"""レーザ材料発注行のキー変更（単位1a）のテスト。

発注行のキーを (material, required_date, supplier) に変更したことに対して、
POST / GET / 入荷実績 / 注文書出力 / 重複統合マイグレーションの動作を確認する。
"""
import importlib
from datetime import date, datetime
from io import BytesIO
from types import SimpleNamespace
from unittest import mock

from django.contrib.auth.models import User
from django.test import SimpleTestCase, TestCase
from openpyxl import load_workbook
from reportlab.pdfgen import canvas as reportlab_canvas
from rest_framework.test import APIRequestFactory, force_authenticate

from accounts.models import ApprovalRequest, ApprovalRouteConfig
from masters.models import Product, Supplier
from production.models_laser_weekly_plan import LaserMaterialReceipt, LaserWeeklyMaterialOrderProgress
from production.views_laser_weekly_plan import LaserWeeklyPlanViewSet


MEISEI = LaserWeeklyMaterialOrderProgress.SUPPLIER_MEISEI
SATO = LaserWeeklyMaterialOrderProgress.SUPPLIER_SATO


class LaserMaterialOrderProgressTestBase(TestCase):
    def setUp(self):
        self.factory = APIRequestFactory()
        self.user = User.objects.create_user(username='material-order-user', password='test-pass')
        # 承認ルートはマイグレーション 0058 で作成済みのため、取得または作成する。
        self.route, _ = ApprovalRouteConfig.objects.get_or_create(
            item_key='laser_material_order',
            defaults={
                'item_name': 'レーザー材料発注',
                'creator_role': 'leader',
                'reviewer1_role': 'supervisor',
                'reviewer2_enabled': False,
                'approver_role': 'manager',
            },
        )
        self.material = Product.objects.create(product_code='SPCC 1.6x1219x2438', product_name='テスト材料A')
        self.material_b = Product.objects.create(product_code='SPHC 2.3x1219x2438', product_name='テスト材料B')

    # --- 補助 ---
    def _post_progress(self, start_date, items):
        request = self.factory.post(
            '/api/laser-weekly-plans/material-order-progress/',
            {'start_date': start_date, 'items': items},
            format='json',
        )
        force_authenticate(request, user=self.user)
        return LaserWeeklyPlanViewSet.as_view({'post': 'material_order_progress'})(request)

    def _get_progress(self, start_date, display_start, display_end):
        request = self.factory.get(
            '/api/laser-weekly-plans/material-order-progress/',
            {'start_date': start_date, 'display_start': display_start, 'display_end': display_end},
        )
        force_authenticate(request, user=self.user)
        return LaserWeeklyPlanViewSet.as_view({'get': 'material_order_progress'})(request)

    def _item(self, required_date, meisei_lots=0, meisei_sheets=0, sato_lots=0, sato_sheets=0, sato_enabled=False, material=None):
        return {
            'material_id': (material or self.material).id,
            'required_date': required_date,
            'delivery_date': required_date,
            'required_sheets': 10,
            'lot_multiple': 5,
            'meisei_lots': meisei_lots,
            'meisei_sheets': meisei_sheets,
            'sato_lots': sato_lots,
            'sato_sheets': sato_sheets,
            'sato_enabled': sato_enabled,
        }

    def _create_approval(self, supplier, start_date, lock_start, lock_end, status='approved', extra_context=None):
        context = {
            'supplier': supplier,
            'start_date': start_date,
            'lock_start_date': lock_start,
            'lock_end_date': lock_end,
            'order_created': True,
        }
        context.update(extra_context or {})
        return ApprovalRequest.objects.create(
            route_config=self.route,
            creator=self.user,
            status=status,
            current_stage='completed' if status in ('approved', 'sent') else 'creator',
            context=context,
        )

    def _create_row(self, supplier, required_date, order_lots, order_sheets=0, plan_start_date=None, delivery_date=None, material=None, is_manual=False):
        return LaserWeeklyMaterialOrderProgress.objects.create(
            plan_start_date=plan_start_date,
            material=material or self.material,
            required_date=None if is_manual else required_date,
            delivery_date=delivery_date or required_date,
            supplier=supplier,
            required_sheets=10,
            lot_multiple=5,
            required_lots=2,
            order_lots=order_lots,
            order_sheets=order_sheets,
            is_manual=is_manual,
        )

    def _planned_rows(self, required_date, supplier):
        return LaserWeeklyMaterialOrderProgress.objects.filter(
            material=self.material, required_date=required_date, supplier=supplier, is_manual=False,
        )


class MaterialOrderProgressPostTest(LaserMaterialOrderProgressTestBase):
    def test_b1_post_from_two_weeks_keeps_one_row_with_latest_values(self):
        """B1: 別の週開始日から同じ (材料, 日, 仕入先) を2回保存しても1行で、値は後の保存の値。"""
        first = self._post_progress('2026-09-07', [self._item('2026-09-16', meisei_lots=2)])
        second = self._post_progress('2026-09-14', [self._item('2026-09-16', meisei_lots=4, meisei_sheets=3)])

        self.assertEqual(first.status_code, 200)
        self.assertEqual(second.status_code, 200)
        rows = list(self._planned_rows(date(2026, 9, 16), MEISEI))
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0].order_lots, 4)
        self.assertEqual(rows[0].order_sheets, 3)
        self.assertEqual(rows[0].plan_start_date, date(2026, 9, 14))

    def test_b2_post_zero_updates_row_and_does_not_delete(self):
        """B2: 手数0を保存すると既存行が0に更新され、削除されない。"""
        self._create_row(MEISEI, date(2026, 9, 16), order_lots=3, order_sheets=2, plan_start_date=date(2026, 9, 7))

        response = self._post_progress('2026-09-14', [self._item('2026-09-16', meisei_lots=0, meisei_sheets=0)])

        self.assertEqual(response.status_code, 200)
        rows = list(self._planned_rows(date(2026, 9, 16), MEISEI))
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0].order_lots, 0)
        self.assertEqual(rows[0].order_sheets, 0)

    def test_b2_post_zero_creates_row_when_missing(self):
        """B2補足: 行がない日に手数0を保存しても、0の行が作られる。"""
        response = self._post_progress('2026-09-14', [self._item('2026-09-17')])

        self.assertEqual(response.status_code, 200)
        rows = list(self._planned_rows(date(2026, 9, 17), MEISEI))
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0].order_lots, 0)

    def test_b3_locked_supplier_date_is_not_written_and_changes_are_returned(self):
        """B3: ロックされた (仕入先, 日) は書かない。値が違う場合だけ skipped_locked_changes に入る。"""
        self._create_approval(MEISEI, '2026-09-07', '2026-09-14', '2026-09-18')
        locked_row = self._create_row(MEISEI, date(2026, 9, 15), order_lots=2, plan_start_date=date(2026, 9, 7))
        locked_row.refresh_from_db()
        before_updated_at = locked_row.updated_at

        response = self._post_progress('2026-09-14', [
            self._item('2026-09-15', meisei_lots=5, sato_lots=3, sato_enabled=True),
        ])

        self.assertEqual(response.status_code, 200)
        locked_row.refresh_from_db()
        self.assertEqual(locked_row.order_lots, 2)
        self.assertEqual(locked_row.updated_at, before_updated_at)
        self.assertEqual(locked_row.plan_start_date, date(2026, 9, 7))
        self.assertEqual(response.data['skipped_locked_count'], 1)
        self.assertEqual(len(response.data['skipped_locked_changes']), 1)
        change = response.data['skipped_locked_changes'][0]
        self.assertEqual(change['material_id'], self.material.id)
        self.assertEqual(change['product_code'], self.material.product_code)
        self.assertEqual(change['required_date'], '2026-09-15')
        self.assertEqual(change['supplier'], MEISEI)
        self.assertEqual(change['sent_order_lots'], 5)
        self.assertEqual(change['db_order_lots'], 2)
        # もう一方の仕入先（ロックなし）は更新される
        sato_rows = list(self._planned_rows(date(2026, 9, 15), SATO))
        self.assertEqual(len(sato_rows), 1)
        self.assertEqual(sato_rows[0].order_lots, 3)

        # 同じ値を送った場合は skipped_locked_changes に入らない（件数には入る）
        same_response = self._post_progress('2026-09-14', [
            self._item('2026-09-15', meisei_lots=2, sato_lots=3, sato_enabled=True),
        ])
        self.assertEqual(same_response.status_code, 200)
        self.assertEqual(same_response.data['skipped_locked_count'], 1)
        self.assertEqual(same_response.data['skipped_locked_changes'], [])

    def test_b3_locked_without_row_returns_change_only_when_nonzero(self):
        """B3補足: ロック日にDBの行がない場合、0以外を送ったときだけ返す。"""
        self._create_approval(MEISEI, '2026-09-07', '2026-09-14', '2026-09-18')

        zero_response = self._post_progress('2026-09-14', [self._item('2026-09-16')])
        nonzero_response = self._post_progress('2026-09-14', [self._item('2026-09-16', meisei_sheets=4)])

        self.assertEqual(zero_response.data['skipped_locked_changes'], [])
        self.assertEqual(len(nonzero_response.data['skipped_locked_changes']), 1)
        self.assertIsNone(nonzero_response.data['skipped_locked_changes'][0]['db_order_lots'])
        self.assertFalse(self._planned_rows(date(2026, 9, 16), MEISEI).exists())

    def test_b3_locked_sato_row_is_not_deleted_when_sato_disabled(self):
        """B3補足: ロックされた日の SATO 行は、sato_enabled が False でも削除しない。"""
        self._create_approval(SATO, '2026-09-07', '2026-09-14', '2026-09-18')
        sato_row = self._create_row(SATO, date(2026, 9, 15), order_lots=1, plan_start_date=date(2026, 9, 7))

        response = self._post_progress('2026-09-14', [self._item('2026-09-15', meisei_lots=1, sato_enabled=False)])

        self.assertEqual(response.status_code, 200)
        self.assertTrue(LaserWeeklyMaterialOrderProgress.objects.filter(id=sato_row.id).exists())

    def test_unlocked_sato_row_is_deleted_when_sato_disabled(self):
        """ロックされていない日の SATO 行は、sato_enabled が False なら新しいキーで削除する。"""
        self._create_row(SATO, date(2026, 9, 15), order_lots=1, plan_start_date=date(2026, 9, 7))

        response = self._post_progress('2026-09-14', [self._item('2026-09-15', meisei_lots=1, sato_enabled=False)])

        self.assertEqual(response.status_code, 200)
        self.assertFalse(self._planned_rows(date(2026, 9, 15), SATO).exists())

    def test_created_without_order_created_does_not_lock(self):
        """created で order_created が True でない承認はロックしない（従来どおり）。"""
        self._create_approval(MEISEI, '2026-09-07', '2026-09-14', '2026-09-18', status='created', extra_context={'order_created': False})

        response = self._post_progress('2026-09-14', [self._item('2026-09-15', meisei_lots=2)])

        self.assertEqual(response.data['skipped_locked_count'], 0)
        self.assertEqual(self._planned_rows(date(2026, 9, 15), MEISEI).get().order_lots, 2)

    def test_b4_adjustment_editing_period_is_writable(self):
        """B4: 納期調整中 (editing=True) の調整期間は POST で更新される。"""
        self._create_approval(
            MEISEI, '2026-09-07', '2026-09-14', '2026-09-18',
            extra_context={'material_order_adjustments': {MEISEI: {
                'lock_start_date': '2026-09-14', 'lock_end_date': '2026-09-18', 'editing': True,
            }}},
        )
        row = self._create_row(MEISEI, date(2026, 9, 15), order_lots=2, plan_start_date=date(2026, 9, 7))

        response = self._post_progress('2026-09-14', [self._item('2026-09-15', meisei_lots=6)])

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data['skipped_locked_count'], 0)
        row.refresh_from_db()
        self.assertEqual(row.order_lots, 6)

    def test_b4_saved_adjustment_locks_again(self):
        """B4補足: 納期調整を保存した後 (editing=False) は再びロックされる。"""
        self._create_approval(
            MEISEI, '2026-09-07', '2026-09-14', '2026-09-18',
            extra_context={'material_order_adjustments': {MEISEI: {
                'lock_start_date': '2026-09-14', 'lock_end_date': '2026-09-18', 'editing': False,
            }}},
        )
        row = self._create_row(MEISEI, date(2026, 9, 15), order_lots=2, plan_start_date=date(2026, 9, 7))

        self._post_progress('2026-09-14', [self._item('2026-09-15', meisei_lots=6)])

        row.refresh_from_db()
        self.assertEqual(row.order_lots, 2)

    def test_b4_other_locking_approval_keeps_lock_during_editing(self):
        """B4補足: 調整中の承認とは別の承認が同じ (仕入先, 日) をロックしていれば、ロックのまま。"""
        self._create_approval(
            MEISEI, '2026-09-07', '2026-09-14', '2026-09-18',
            extra_context={'material_order_adjustments': {MEISEI: {
                'lock_start_date': '2026-09-14', 'lock_end_date': '2026-09-18', 'editing': True,
            }}},
        )
        self._create_approval(MEISEI, '2026-09-14', '2026-09-15', '2026-09-16')
        row = self._create_row(MEISEI, date(2026, 9, 15), order_lots=2, plan_start_date=date(2026, 9, 7))

        self._post_progress('2026-09-14', [self._item('2026-09-15', meisei_lots=6)])

        row.refresh_from_db()
        self.assertEqual(row.order_lots, 2)

    def test_invalid_item_saves_nothing(self):
        """不正値が1件でもあれば何も保存しない。"""
        bad = self._item('2026-09-16', meisei_lots=1)
        bad['lot_multiple'] = 0

        response = self._post_progress('2026-09-14', [self._item('2026-09-15', meisei_lots=1), bad])

        self.assertEqual(response.status_code, 400)
        self.assertFalse(LaserWeeklyMaterialOrderProgress.objects.exists())


class MaterialOrderProgressGetTest(LaserMaterialOrderProgressTestBase):
    def test_b5_get_returns_rows_in_display_period_without_duplicate_keys(self):
        """B5: items は表示期間内の required_date の行だけで、キーの重複がない。"""
        self._create_row(MEISEI, date(2026, 9, 7), order_lots=1, plan_start_date=date(2026, 8, 31))
        self._create_row(MEISEI, date(2026, 9, 15), order_lots=2, plan_start_date=date(2026, 9, 7))
        self._create_row(SATO, date(2026, 9, 15), order_lots=0, plan_start_date=date(2026, 9, 14))
        self._create_row(MEISEI, date(2026, 9, 25), order_lots=3, plan_start_date=date(2026, 9, 21), material=self.material_b)
        self._create_row(MEISEI, date(2026, 9, 28), order_lots=4, plan_start_date=date(2026, 9, 28))
        self._create_row(MEISEI, None, order_lots=5, delivery_date=date(2026, 9, 16), is_manual=True)

        response = self._get_progress('2026-09-14', '2026-09-07', '2026-09-27')

        self.assertEqual(response.status_code, 200)
        keys = [(item['material_id'], item['required_date'], item['supplier']) for item in response.data['items']]
        self.assertEqual(len(keys), len(set(keys)))
        self.assertEqual(set(keys), {
            (self.material.id, '2026-09-07', MEISEI),
            (self.material.id, '2026-09-15', MEISEI),
            (self.material.id, '2026-09-15', SATO),
            (self.material_b.id, '2026-09-25', MEISEI),
        })
        self.assertEqual(response.data['overlapping_items'], [])
        self.assertEqual(response.data['locked_items'], [])


class MaterialReceiptsGetTest(LaserMaterialOrderProgressTestBase):
    def test_b6_material_receipts_returns_each_key_once(self):
        """B6: 入荷実績一覧は同じ (材料, 仕入先, 必要日) を2回返さない。"""
        order = self._create_row(MEISEI, date(2026, 9, 15), order_lots=2, plan_start_date=date(2026, 9, 7))
        self._create_row(SATO, date(2026, 9, 15), order_lots=1, plan_start_date=date(2026, 9, 14))
        self._create_row(MEISEI, date(2026, 9, 16), order_lots=0, plan_start_date=date(2026, 9, 14))
        LaserMaterialReceipt.objects.create(order=order, received_date=date(2026, 9, 15), received_lots=1, received_by=self.user)
        self._create_approval(MEISEI, '2026-09-07', '2026-09-14', '2026-09-18')

        request = self.factory.get(
            '/api/laser-weekly-plans/material-receipts/',
            {'start_date': '2026-09-14', 'end_date': '2026-09-20'},
        )
        force_authenticate(request, user=self.user)
        response = LaserWeeklyPlanViewSet.as_view({'get': 'material_receipts'})(request)

        self.assertEqual(response.status_code, 200)
        keys = [(item['material_id'], item['supplier'], item['delivery_date']) for item in response.data['items']]
        self.assertEqual(len(keys), len(set(keys)))
        self.assertEqual(len(keys), 2)  # 0の行は出さない
        by_supplier = {item['supplier']: item for item in response.data['items']}
        self.assertEqual(len(by_supplier[MEISEI]['receipts']), 1)
        self.assertEqual(by_supplier[MEISEI]['source_status'], 'ORDER_LOCKED')
        self.assertEqual(by_supplier[SATO]['source_status'], 'PLAN_SAVED')


class _RecordingCanvas(reportlab_canvas.Canvas):
    """PDFに描いた文字列を記録するテスト用キャンバス。"""
    drawn_texts = []

    def drawString(self, x, y, text, *args, **kwargs):
        _RecordingCanvas.drawn_texts.append(str(text))
        return super().drawString(x, y, text, *args, **kwargs)

    def drawCentredString(self, x, y, text, *args, **kwargs):
        _RecordingCanvas.drawn_texts.append(str(text))
        return super().drawCentredString(x, y, text, *args, **kwargs)

    def drawRightString(self, x, y, text, *args, **kwargs):
        _RecordingCanvas.drawn_texts.append(str(text))
        return super().drawRightString(x, y, text, *args, **kwargs)


class MaterialOrderExportTest(LaserMaterialOrderProgressTestBase):
    def setUp(self):
        super().setUp()
        Supplier.objects.get_or_create(
            supplier_code='000048',
            defaults={'supplier_name': '名成鋼機テスト', 'supplier_type': 'both'},
        )
        # 週開始日の違う行。どちらも期間内なので全部出力される。
        self._create_row(MEISEI, date(2026, 9, 15), order_lots=2, plan_start_date=date(2026, 9, 7))
        self._create_row(MEISEI, date(2026, 9, 17), order_lots=3, plan_start_date=date(2026, 9, 14), material=self.material_b)
        # 出力されない行: 0の行、他の仕入先、期間外
        self._create_row(MEISEI, date(2026, 9, 16), order_lots=0, plan_start_date=date(2026, 9, 14))
        self._create_row(SATO, date(2026, 9, 16), order_lots=7, plan_start_date=date(2026, 9, 14))
        self._create_row(MEISEI, date(2026, 9, 21), order_lots=9, plan_start_date=date(2026, 9, 14))

    def test_b7_excel_contains_all_rows_in_period_regardless_of_week(self):
        """B7: Excel の抽出は週開始日に関係なく、期間内の行を全部含む。"""
        request = self.factory.get(
            '/api/laser-weekly-plans/material-order-excel/',
            {'plan_start_date': '2026-09-14', 'start_date': '2026-09-14', 'end_date': '2026-09-18', 'supplier': MEISEI},
        )
        force_authenticate(request, user=self.user)
        response = LaserWeeklyPlanViewSet.as_view({'get': 'material_order_excel'})(request)

        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(b''.join(response.streaming_content)))
        worksheet = workbook.active
        materials = [worksheet.cell(row=row, column=3).value for row in range(8, 10)]
        self.assertEqual(materials, ['SPCC', 'SPHC'])
        # 合計列（6列目）: 材料A=2, 材料B=3
        self.assertEqual(worksheet.cell(row=8, column=6).value, '2')
        self.assertEqual(worksheet.cell(row=9, column=6).value, '3')
        # 材料行の次は空行（他の仕入先・期間外・0の行は出ない）
        self.assertIsNone(worksheet.cell(row=10, column=3).value)

    def test_b7_pdf_contains_all_rows_in_period_regardless_of_week(self):
        """B7: PDF の抽出は週開始日に関係なく、期間内の行を全部含む。"""
        self._create_approval(MEISEI, '2026-09-14', '2026-09-14', '2026-09-18', status='created')
        _RecordingCanvas.drawn_texts = []

        with mock.patch.object(reportlab_canvas, 'Canvas', _RecordingCanvas):
            pdf_bytes, _filename = LaserWeeklyPlanViewSet()._build_material_order_pdf(
                date(2026, 9, 14), date(2026, 9, 14), date(2026, 9, 18), MEISEI,
            )

        self.assertTrue(pdf_bytes.startswith(b'%PDF'))
        texts = _RecordingCanvas.drawn_texts
        self.assertIn('SPCC', texts)
        self.assertIn('SPHC', texts)
        self.assertNotIn('7', texts)  # SATO の行（7ロット）は出ない
        self.assertNotIn('9', texts)  # 期間外の行（9ロット）は出ない


# --- 重複統合マイグレーションの関数 ---
# 新しい一意索引があるテストDB(MySQL)には重複行を作れないため、メモリ上の疑似モデルで検証する。

merge_migration = importlib.import_module('production.migrations.0116_merge_laser_material_order_duplicate_rows')


class _FakeQuerySet:
    def __init__(self, manager, rows):
        self.manager = manager
        self.rows = list(rows)

    @staticmethod
    def _value(obj, path):
        for part in path.split('__'):
            obj = getattr(obj, part)
        return obj

    def filter(self, **lookups):
        rows = self.rows
        for key, expected in lookups.items():
            if key.endswith('__in'):
                field = key[:-4]
                rows = [r for r in rows if self._value(r, field) in expected]
            elif key.endswith('__isnull'):
                field = key[:-8]
                rows = [r for r in rows if (self._value(r, field) is None) == expected]
            else:
                rows = [r for r in rows if self._value(r, key) == expected]
        return _FakeQuerySet(self.manager, rows)

    def order_by(self, field):
        return _FakeQuerySet(self.manager, sorted(self.rows, key=lambda r: getattr(r, field)))

    def values_list(self, field, flat=False):
        return [getattr(r, field) for r in self.rows]

    def delete(self):
        ids = {r.id for r in self.rows}
        self.manager.deleted_ids.update(ids)
        self.manager.rows = [r for r in self.manager.rows if r.id not in ids]

    def __iter__(self):
        return iter(self.rows)


class _FakeManager:
    def __init__(self, rows):
        self.rows = list(rows)
        self.deleted_ids = set()

    def filter(self, **lookups):
        return _FakeQuerySet(self, self.rows).filter(**lookups)


def _fake_model(rows):
    return SimpleNamespace(objects=_FakeManager(rows))


def _fake_order(row_id, plan_start_date, required_date=date(2026, 9, 15), supplier=MEISEI, material_id=1, is_manual=False, updated_at=None):
    return SimpleNamespace(
        id=row_id, plan_start_date=plan_start_date, required_date=required_date, supplier=supplier,
        material_id=material_id, is_manual=is_manual, updated_at=updated_at or datetime(2026, 9, 1, 10, 0),
    )


def _fake_approval(supplier, start_date, lock_start, lock_end, status='approved', order_created=True):
    return SimpleNamespace(
        status=status,
        route_config=SimpleNamespace(item_key='laser_material_order'),
        context={
            'supplier': supplier, 'start_date': start_date,
            'lock_start_date': lock_start, 'lock_end_date': lock_end, 'order_created': order_created,
        },
    )


class MergeDuplicateOrderRowsTest(SimpleTestCase):
    def test_keeps_row_of_locked_week(self):
        """ロックされた週の行が1行なら、その行を残す（新しい行があっても）。"""
        orders = _fake_model([
            _fake_order(1, date(2026, 9, 7)),
            _fake_order(2, date(2026, 9, 14), updated_at=datetime(2026, 9, 10, 9, 0)),
            _fake_order(3, date(2026, 9, 1)),
        ])
        receipts = _fake_model([])
        approvals = _fake_model([_fake_approval(MEISEI, '2026-09-07', '2026-09-14', '2026-09-18')])

        result = merge_migration.merge_duplicate_order_rows(orders, receipts, approvals)

        self.assertEqual(result['group_count'], 1)
        self.assertEqual([r.id for r in orders.objects.rows], [1])
        self.assertEqual(orders.objects.deleted_ids, {2, 3})

    def test_keeps_latest_row_when_no_locked_week(self):
        """ロックされた週の行がなければ (plan_start_date, updated_at, id) が最大の行を残す。"""
        orders = _fake_model([
            _fake_order(1, date(2026, 9, 14), updated_at=datetime(2026, 9, 10, 9, 0)),
            _fake_order(2, date(2026, 9, 14), updated_at=datetime(2026, 9, 11, 9, 0)),
            _fake_order(3, date(2026, 9, 7), updated_at=datetime(2026, 9, 12, 9, 0)),
            _fake_order(4, None, required_date=None, is_manual=True),
        ])

        merge_migration.merge_duplicate_order_rows(orders, _fake_model([]), _fake_model([]))

        self.assertEqual(sorted(r.id for r in orders.objects.rows), [2, 4])

    def test_raises_when_two_rows_are_in_locked_weeks(self):
        """ロックされた週の行が2行以上ある組があれば例外で止め、何も消さない。"""
        orders = _fake_model([
            _fake_order(1, date(2026, 9, 7)),
            _fake_order(2, date(2026, 9, 14)),
        ])
        approvals = _fake_model([
            _fake_approval(MEISEI, '2026-09-07', '2026-09-14', '2026-09-18'),
            _fake_approval(MEISEI, '2026-09-14', '2026-09-15', '2026-09-16'),
        ])

        with self.assertRaises(RuntimeError) as ctx:
            merge_migration.merge_duplicate_order_rows(orders, _fake_model([]), approvals)

        self.assertIn('required_date=2026-09-15', str(ctx.exception))
        self.assertEqual(orders.objects.deleted_ids, set())

    def test_raises_when_deleted_row_has_receipt(self):
        """消す行に入荷実績がつながっていたら例外で止め、何も消さない。"""
        orders = _fake_model([
            _fake_order(1, date(2026, 9, 7)),
            _fake_order(2, date(2026, 9, 14)),
        ])
        receipts = _fake_model([SimpleNamespace(id=10, order_id=1)])

        with self.assertRaises(RuntimeError) as ctx:
            merge_migration.merge_duplicate_order_rows(orders, receipts, _fake_model([]))

        self.assertIn('order_id=1', str(ctx.exception))
        self.assertEqual(orders.objects.deleted_ids, set())

    def test_created_without_order_created_is_not_locked(self):
        """created で order_created が True でない承認はロックとして扱わない。"""
        orders = _fake_model([
            _fake_order(1, date(2026, 9, 7)),
            _fake_order(2, date(2026, 9, 14)),
        ])
        approvals = _fake_model([
            _fake_approval(MEISEI, '2026-09-07', '2026-09-14', '2026-09-18', status='created', order_created=False),
        ])

        merge_migration.merge_duplicate_order_rows(orders, _fake_model([]), approvals)

        self.assertEqual([r.id for r in orders.objects.rows], [2])
