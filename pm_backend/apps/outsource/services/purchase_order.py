from io import BytesIO
from collections import defaultdict
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from outsource.models import MaterialRequirement


def generate_purchase_orders(material_ids=None, unsupplied_only=True, unordered_only=False):
    """
    材料所要量からメーカ別注文書Excelを生成する。
    メーカ（調達先）ごとにシートを分けて出力。
    """
    qs = MaterialRequirement.objects.select_related('split', 'split__order')
    if material_ids:
        qs = qs.filter(id__in=material_ids)
    if unordered_only:
        qs = qs.filter(ordered=False)
    elif unsupplied_only:
        qs = qs.filter(supplied=False)
    qs = qs.order_by('supplier_name', 'supply_date', 'material_code')

    # メーカ別にグループ化
    by_supplier = defaultdict(list)
    for m in qs:
        if float(m.order_qty or 0) <= 0:
            continue
        supplier = m.supplier_name or '調達先未設定'
        by_supplier[supplier].append(m)

    if not by_supplier:
        return None, []

    wb = Workbook()
    wb.remove(wb.active)

    # スタイル
    title_font = Font(bold=True, size=14)
    header_font = Font(bold=True, size=10)
    header_fill = PatternFill('solid', fgColor='D9E1F2')
    thin_border = Border(
        left=Side('thin'), right=Side('thin'),
        top=Side('thin'), bottom=Side('thin'),
    )
    right_align = Alignment(horizontal='right')
    center_align = Alignment(horizontal='center')

    supplier_list = []

    for supplier, items in by_supplier.items():
        sheet_name = supplier[:31]  # Excelシート名は31文字制限
        ws = wb.create_sheet(title=sheet_name)

        # ヘッダー
        ws['A1'] = '注 文 書'
        ws['A1'].font = title_font

        ws['A3'] = f'{supplier} 御中'
        ws['A3'].font = Font(bold=True, size=12)

        ws['A4'] = f'発行日: {date.today().strftime("%Y/%m/%d")}'
        ws['A5'] = f'発行元: ダイソウ工業株式会社'

        # 材料を材料コード+支給日で集約
        agg = defaultdict(lambda: {
            'material_code': '', 'material_name': '',
            'total_qty': 0, 'supply_date': None, 'cases': set()
        })
        for m in items:
            key = (m.material_code, m.supply_date)
            row = agg[key]
            row['material_code'] = m.material_code
            row['material_name'] = m.material_name
            row['total_qty'] += float(m.order_qty or 0)
            row['supply_date'] = m.supply_date
            row['cases'].add(m.split.order.case_no)

        # テーブル
        headers = ['No', '材料コード', '材料名称', '数量', '納入希望日', '関連案件']
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=7, column=col, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.border = thin_border
            cell.alignment = center_align

        total_qty = 0
        for idx, ((mat_code, sup_date), data) in enumerate(
            sorted(agg.items(), key=lambda x: (x[0][1], x[0][0])), 1
        ):
            row_num = 7 + idx
            vals = [
                idx,
                data['material_code'],
                data['material_name'],
                data['total_qty'],
                data['supply_date'],
                ', '.join(sorted(data['cases'])),
            ]
            for col, val in enumerate(vals, 1):
                cell = ws.cell(row=row_num, column=col, value=val)
                cell.border = thin_border
                if col == 1:
                    cell.alignment = center_align
                elif col == 4:
                    cell.alignment = right_align
                    cell.number_format = '#,##0'
                elif col == 5 and isinstance(val, date):
                    cell.number_format = 'YYYY/MM/DD'
                    cell.alignment = center_align

            total_qty += data['total_qty']

        # 合計行
        total_row = 7 + len(agg) + 1
        ws.cell(row=total_row, column=3, value='合計').font = Font(bold=True)
        ws.cell(row=total_row, column=3).border = thin_border
        total_cell = ws.cell(row=total_row, column=4, value=total_qty)
        total_cell.font = Font(bold=True)
        total_cell.alignment = right_align
        total_cell.number_format = '#,##0'
        total_cell.border = thin_border

        # 備考欄
        note_row = total_row + 2
        ws.cell(row=note_row, column=1, value='備考:')
        ws.cell(row=note_row + 1, column=1, value='・納入希望日は材料コードごとに異なる場合があります。')

        # 列幅
        widths = [5, 18, 28, 10, 14, 30]
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[ws.cell(row=7, column=i).column_letter].width = w

        supplier_list.append({
            'supplier': supplier,
            'item_count': len(agg),
            'material_count': len(items),
        })

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output, supplier_list
