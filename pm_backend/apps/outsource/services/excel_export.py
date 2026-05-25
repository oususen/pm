from io import BytesIO
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from outsource.models import OutsourceOrder


def generate_split_plan_excel(order_ids):
    """
    外作先へ送付する分割計画テンプレートExcelを生成する。
    制約条件（最早着手日・最遅完了日）を記載し、
    外作先が加工日・数量を記入して返却する。
    """
    orders = OutsourceOrder.objects.filter(
        id__in=order_ids
    ).select_related('item', 'item__subcontractor').order_by('painting_date', 'item_code')

    wb = Workbook()
    ws = wb.active
    ws.title = '分割計画'

    # スタイル定義
    header_font = Font(bold=True, size=10)
    header_fill = PatternFill('solid', fgColor='D9E1F2')
    input_fill = PatternFill('solid', fgColor='FFFFCC')
    thin_border = Border(
        left=Side('thin'), right=Side('thin'),
        top=Side('thin'), bottom=Side('thin')
    )

    # ヘッダー行
    headers = [
        '案件番号', '品目コード', '品番', '品目名称', '受注数量',
        '塗装名', '塗装日', '最早着手日', '最遅完了日',
        '加工日1', '数量1', '加工日2', '数量2',
        '加工日3', '数量3', '加工日4', '数量4',
        '加工日5', '数量5',
    ]

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = thin_border
        cell.alignment = Alignment(horizontal='center', wrap_text=True)

    # データ行
    for row_idx, order in enumerate(orders, 2):
        # 制約条件を計算（未計算の場合）
        if order.item and (not order.earliest_start or not order.latest_finish):
            order.calculate_constraints()
            order.save()

        data = [
            order.case_no,
            order.item_code,
            (order.item.product_number if order.item else ''),
            order.item_name,
            order.order_qty,
            order.painting_name,
            order.painting_date,
            order.earliest_start,
            order.latest_finish,
        ]

        for col, val in enumerate(data, 1):
            cell = ws.cell(row=row_idx, column=col, value=val)
            cell.border = thin_border
            cell.alignment = Alignment(horizontal='center' if col >= 4 else 'left')
            if isinstance(val, date):
                cell.number_format = 'YYYY/MM/DD'

        # 記入欄（黄色背景）: 加工日1〜5, 数量1〜5
        process_start_col = 10
        for col in range(process_start_col, process_start_col + 10):
            cell = ws.cell(row=row_idx, column=col)
            cell.fill = input_fill
            cell.border = thin_border
            if (col - process_start_col) % 2 == 1:  # 数量列（加工日の次列）
                cell.number_format = '#,##0'

    # 列幅調整
    col_widths = [22, 16, 14, 24, 8, 14, 12, 12, 12, 12, 8, 12, 8, 12, 8, 12, 8, 12, 8]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w

    # 注意書きシート
    ws2 = wb.create_sheet('注意事項')
    ws2['A1'] = '【分割計画記入ルール】'
    ws2['A1'].font = Font(bold=True, size=12)
    ws2['A3'] = '1. 加工日は「最早着手日」〜「最遅完了日」の範囲内で記入してください。'
    ws2['A4'] = '2. 数量の合計が「受注数量」と一致するようにしてください。'
    ws2['A5'] = '3. 加工日は YYYY/MM/DD 形式で記入してください。'
    ws2['A6'] = '4. 分割が5回を超える場合は列を追加してください。'
    ws2.column_dimensions['A'].width = 60

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output
