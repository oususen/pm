import re
from datetime import date, datetime

from openpyxl import load_workbook

from outsource.models import OutsourceOrder, OutsourceSplit


def _parse_date(val, reference_date=None):
    """
    様々な日付形式をパースする。
    対応形式:
      - datetime / date オブジェクト（Excel日付セル）
      - int / float（Excelシリアル日付）
      - '2026/06/01', '2026-06-01'
      - '6月1日', '6月01日'（年はreference_dateから推定）
      - '6/1', '06/01'（年はreference_dateから推定）
    """
    if isinstance(val, datetime):
        return val.date()
    if isinstance(val, date):
        return val

    # Excelシリアル日付（int/float）
    if isinstance(val, (int, float)) and val > 40000:
        from openpyxl.utils.datetime import from_excel
        return from_excel(val).date()

    if not isinstance(val, str):
        # 最後の手段: str変換して再パース
        val = str(val).strip()
        if not val:
            return None
    else:
        val = val.strip()
        if not val:
            return None

    # YYYY/MM/DD or YYYY-MM-DD
    for fmt in ('%Y/%m/%d', '%Y-%m-%d'):
        try:
            return datetime.strptime(val, fmt).date()
        except ValueError:
            pass

    # 年を推定
    ref_year = reference_date.year if reference_date else date.today().year

    # 「6月1日」「06月01日」
    m = re.match(r'(\d{1,2})月(\d{1,2})日?', val)
    if m:
        return date(ref_year, int(m.group(1)), int(m.group(2)))

    # 「6/1」「06/01」
    m = re.match(r'^(\d{1,2})/(\d{1,2})$', val)
    if m:
        return date(ref_year, int(m.group(1)), int(m.group(2)))

    return None


def import_split_plan_excel(file_content, dry_run=False):
    """
    外作先が記入した分割計画Excelを取り込む。

    Excel構成:
        案件番号 | 品目コード | 品目名称 | 受注数量 | 塗装名 | 塗装日 |
        最早着手日 | 最遅完了日 | 加工日1 | 数量1 | 加工日2 | 数量2 | ...

    dry_run=True: パース・バリデーションのみ（DB変更なし）
    戻り値: { 'updated': [...], 'errors': [...], 'warnings': [], 'has_overwrite': bool }
    """
    results = {'updated': [], 'errors': [], 'warnings': [], 'has_overwrite': False}

    wb = load_workbook(file_content, data_only=True)
    ws = wb.active

    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not row or not row[0]:
            continue

        case_no = str(row[0]).strip()

        try:
            order = OutsourceOrder.objects.get(case_no=case_no)
        except OutsourceOrder.DoesNotExist:
            results['errors'].append({
                'row': row_idx,
                'case_no': case_no,
                'message': f'案件番号 {case_no} が見つかりません',
            })
            continue

        # 加工日・数量のペアを抽出（col 8以降: 加工日1, 数量1, 加工日2, 数量2, ...）
        splits_data = []
        print(f'[DEBUG] row {row_idx}: case={case_no}, total_cols={len(row)}, raw_data={list(row[8:])}')
        col_idx = 8  # 0-indexed: col 8 = 加工日1
        while col_idx < len(row) - 1:
            process_date_val = row[col_idx]
            qty_val = row[col_idx + 1]
            col_idx += 2

            if not process_date_val or not qty_val:
                continue

            process_date = _parse_date(process_date_val, reference_date=order.painting_date)
            if not process_date:
                continue

            qty = int(float(qty_val))
            if qty <= 0:
                continue

            splits_data.append({'process_date': process_date, 'qty': qty})

        if not splits_data:
            results['errors'].append({
                'row': row_idx,
                'case_no': case_no,
                'message': '分割計画データがありません',
            })
            continue

        # 数量合計チェック
        total_qty = sum(s['qty'] for s in splits_data)
        if total_qty != order.order_qty:
            results['warnings'].append({
                'row': row_idx,
                'case_no': case_no,
                'message': f'分割合計({total_qty}) ≠ 受注数量({order.order_qty})',
            })

        # 制約チェック
        for s in splits_data:
            if order.earliest_start and s['process_date'] < order.earliest_start:
                results['warnings'].append({
                    'row': row_idx,
                    'case_no': case_no,
                    'message': f"加工日 {s['process_date']} が最早着手日 {order.earliest_start} より前です",
                })
            if order.latest_finish and s['process_date'] > order.latest_finish:
                results['warnings'].append({
                    'row': row_idx,
                    'case_no': case_no,
                    'message': f"加工日 {s['process_date']} が最遅完了日 {order.latest_finish} より後です",
                })

        # 既存分割がある場合は上書き警告
        existing_count = order.splits.count()
        if existing_count > 0:
            results['has_overwrite'] = True
            results['warnings'].append({
                'row': row_idx,
                'case_no': case_no,
                'message': f'既存の分割計画({existing_count}件)があります。上書きしますか？',
            })

        if not dry_run:
            order.splits.all().delete()
            for seq, s in enumerate(splits_data, 1):
                OutsourceSplit.objects.create(
                    order=order,
                    sequence=seq,
                    process_date=s['process_date'],
                    qty=s['qty'],
                )
            order.status = 'SPLIT_REGISTERED'
            order.save()

        results['updated'].append({
            'row': row_idx,
            'case_no': case_no,
            'split_count': len(splits_data),
            'total_qty': total_qty,
        })

    return results
