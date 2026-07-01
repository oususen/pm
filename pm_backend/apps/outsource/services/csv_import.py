import csv
import io
from datetime import date, datetime, timedelta

from openpyxl import load_workbook

from outsource.models import OutsourceOrder, OutsourceItem


FIRST_ARTICLE_CONTACT_TYPES = ['FB初物検査', '初物検査']
FIRST_ARTICLE_LOOKBACK_DAYS = 90


def _decode_csv_text(file_content, encoding):
    if not isinstance(file_content, bytes):
        return file_content
    if file_content.startswith(b'\xef\xbb\xbf'):
        return file_content.decode('utf-8-sig')
    try:
        return file_content.decode(encoding)
    except UnicodeDecodeError:
        return file_content.decode('utf-8-sig')


def _parse_fb_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or '').strip()
    if not text:
        raise ValueError('日付が空です')
    for fmt in ('%Y/%m/%d', '%Y-%m-%d', '%Y%m%d'):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise ValueError(f'日付形式が不正です: {text}')


def _build_painting_name(*parts):
    values = [str(part).strip() for part in parts if str(part or '').strip()]
    return ' / '.join(values)


def _derive_product_number(item_code):
    code = str(item_code or '').strip().lstrip('B')
    if len(code) >= 10:
        return f'{code[:6]}-{code[6:10]}'
    return str(item_code or '').strip()


def _has_recent_order(item_code, painting_date, import_started_at):
    window_start = painting_date - timedelta(days=FIRST_ARTICLE_LOOKBACK_DAYS)
    return OutsourceOrder.objects.filter(
        item_code=item_code,
        painting_date__gte=window_start,
        painting_date__lt=painting_date,
        imported_at__lt=import_started_at,
    ).exists()


def _collect_first_article_candidate(item_code, item_name, painting_date, order_qty, case_no, import_started_at):
    if _has_recent_order(item_code, painting_date, import_started_at):
        return None
    return {
        'item_code': item_code,
        'product_number': _derive_product_number(item_code),
        'item_name': item_name,
        'painting_date': painting_date.isoformat(),
        'order_qty': order_qty,
        'case_no': case_no,
        'lookback_days': FIRST_ARTICLE_LOOKBACK_DAYS,
    }


def build_first_article_email_body(candidates):
    lines = [
        'フランスベッド受注取込で、90日以上受注のなかった品番が検出されました。',
        '初物検査の要否を確認してください。',
        '',
        f'判定条件: 塗装日から{FIRST_ARTICLE_LOOKBACK_DAYS}日遡った期間に同一品番の受注がないこと',
        '',
        '対象一覧:',
    ]
    for item in candidates:
        lines.append(
            (
                f"- 品番: {item['product_number']} / 品名: {item['item_name']} / 塗装日: {item['painting_date']} / "
                f"数量: {item['order_qty']} / 案件番号: {item['case_no']}"
            )
        )
    lines.extend(['', '以上'])
    return '\n'.join(lines)


def _create_order(row_num, item_code, item_name, painting_name, painting_date_value, qty_value, results, raw_data, import_started_at):
    item_code = str(item_code or '').strip()
    item_name = str(item_name or '').strip()
    painting_name = str(painting_name or '').strip()
    qty_text = str(qty_value or '').strip()

    if not all([item_code, item_name, painting_date_value, qty_text]):
        results['errors'].append({
            'row': row_num,
            'message': '必須項目が不足しています',
            'data': raw_data,
        })
        return

    painting_date = _parse_fb_date(painting_date_value)
    order_qty = int(float(qty_text))
    case_no = f"{painting_date.strftime('%Y%m%d')}-{item_code}"

    if OutsourceOrder.objects.filter(case_no=case_no).exists():
        results['skipped'].append({
            'row': row_num,
            'case_no': case_no,
            'message': '既に登録済み',
        })
        return

    item = OutsourceItem.objects.filter(item_code=item_code).first()
    OutsourceOrder.objects.create(
        case_no=case_no,
        item=item,
        item_code=item_code,
        item_name=item_name,
        order_qty=order_qty,
        painting_name=painting_name,
        painting_date=painting_date,
        status='IMPORTED',
    )

    results['created'].append({
        'row': row_num,
        'case_no': case_no,
        'item_code': item_code,
        'item_name': item_name,
        'qty': order_qty,
    })
    candidate = _collect_first_article_candidate(
        item_code=item_code,
        item_name=item_name,
        painting_date=painting_date,
        order_qty=order_qty,
        case_no=case_no,
        import_started_at=import_started_at,
    )
    if candidate:
        results.setdefault('first_article_candidates', []).append(candidate)


def import_fb_csv(file_content, encoding='utf-8-sig', user_id=None):
    """
    FB受注CSVを取り込み、案件（OutsourceOrder）を生成する。

    CSVフォーマット:
        品目コード,品目名称,塗装名,塗装日,数量
    """
    results = {'created': [], 'skipped': [], 'errors': [], 'first_article_candidates': []}
    import_started_at = datetime.now()
    text = _decode_csv_text(file_content, encoding)
    reader = csv.DictReader(io.StringIO(text))

    for row_num, row in enumerate(reader, start=2):
        try:
            _create_order(
                row_num=row_num,
                item_code=row.get('品目コード', ''),
                item_name=row.get('品目名称', ''),
                painting_name=row.get('塗装名', ''),
                painting_date_value=row.get('塗装日', ''),
                qty_value=row.get('数量', ''),
                results=results,
                raw_data=row,
                import_started_at=import_started_at,
            )
        except Exception as e:
            results['errors'].append({
                'row': row_num,
                'message': f'{type(e).__name__}: {e}',
                'data': row,
            })

    return results


def import_fb_excel(file_content, user_id=None):
    """
    FB受注Excelを取り込み、案件（OutsourceOrder）を生成する。

    Excelフォーマット:
        伝票区分,伝票タイプ,品目コード,品目名称,発注数,納入期日
    """
    results = {'created': [], 'skipped': [], 'errors': [], 'first_article_candidates': []}
    import_started_at = datetime.now()
    wb = load_workbook(io.BytesIO(file_content), data_only=True)
    ws = wb[wb.sheetnames[0]]

    last_denpyo_kubun = ''
    last_denpyo_type = ''

    for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not row or not any(cell not in (None, '') for cell in row[:6]):
            continue

        denpyo_kubun = str(row[0] or '').strip() or last_denpyo_kubun
        denpyo_type = str(row[1] or '').strip() or last_denpyo_type
        item_code = row[2] if len(row) > 2 else ''
        item_name = row[3] if len(row) > 3 else ''
        order_qty = row[4] if len(row) > 4 else ''
        nouki = row[5] if len(row) > 5 else ''

        if denpyo_kubun:
            last_denpyo_kubun = denpyo_kubun
        if denpyo_type:
            last_denpyo_type = denpyo_type

        raw_data = {
            '伝票区分': denpyo_kubun,
            '伝票タイプ': denpyo_type,
            '品目コード': item_code,
            '品目名称': item_name,
            '発注数': order_qty,
            '納入期日': nouki,
        }

        try:
            _create_order(
                row_num=row_num,
                item_code=item_code,
                item_name=item_name,
                painting_name=_build_painting_name(denpyo_type, denpyo_kubun),
                painting_date_value=nouki,
                qty_value=order_qty,
                results=results,
                raw_data=raw_data,
                import_started_at=import_started_at,
            )
        except Exception as e:
            results['errors'].append({
                'row': row_num,
                'message': f'{type(e).__name__}: {e}',
                'data': raw_data,
            })

    return results


def import_fb_order_file(file_content, filename='', encoding='utf-8-sig', user_id=None):
    lower_name = str(filename or '').lower()
    if lower_name.endswith(('.xlsx', '.xlsm', '.xls')):
        return import_fb_excel(file_content, user_id=user_id)
    return import_fb_csv(file_content, encoding=encoding, user_id=user_id)
