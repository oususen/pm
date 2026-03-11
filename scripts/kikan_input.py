"""
kikan_input.py
生産実績照会 Excel出力2 → 基幹システム(SSE0040)自動入力スクリプト

使用方法:
    python kikan_input.py <Excelファイルパス>

例:
    python kikan_input.py production_record2_20260310_20260312.xlsx

事前準備:
    pip install pyautogui pyperclip openpyxl pywin32

注意:
    マウスを画面の左上角に移動すると緊急停止します（pyautoguiフェイルセーフ）
"""

import pyautogui
import pyperclip
import openpyxl
import time
import sys
from pathlib import Path
import win32gui
import win32con

# ============================================================
# ★★ 設定 ★★  実際の画面のTab順に合わせて数値を調整してください
# ============================================================

WINDOW_TITLE_PART = "SSE0040"       # 基幹システムのウィンドウタイトルの一部

# --- Tab数設定（処理区分を起点とするTab回数） ---
# 実際の画面でTabキーを押しながら数えて確認してください
TABS_TO_SEISANBI       = 1   # 処理区分 → 生産日
TABS_TO_HINBAN         = 1   # 生産日   → 品番
TABS_AFTER_HINBAN      = 2   # 品番確定後（品名スキップ等）→ 工程順位  ★要確認
TABS_TO_SEISANSU       = 5   # 工程順位 → 生産数(完成)              ★要確認
TABS_TO_JIKAN_START    = 5   # 生産数(完成) → 加工時間1(開始)       ★要確認
TABS_TO_JIKAN_END      = 1   # 加工時間1(開始) → 加工時間1(終了)    ★要確認

# --- 待機時間（秒） ---
DELAY_KEY       = 0.15   # キー操作間
DELAY_HINBAN    = 0.8    # 品番入力後の品名表示待機
DELAY_REGISTER  = 1.2    # F12登録後の次レコード待機
DELAY_START     = 5.0    # スクリプト開始前のカウントダウン

# ============================================================
# ウィンドウ操作
# ============================================================

def find_window(title_part):
    """部分タイトルでウィンドウハンドルを検索"""
    handles = []
    def callback(hwnd, _):
        if win32gui.IsWindowVisible(hwnd):
            if title_part in win32gui.GetWindowText(hwnd):
                handles.append(hwnd)
    win32gui.EnumWindows(callback, None)
    return handles[0] if handles else None


def activate_window(hwnd):
    """ウィンドウをフォアグラウンドに持ってくる"""
    win32gui.ShowWindow(hwnd, win32con.SW_RESTORE)
    win32gui.SetForegroundWindow(hwnd)
    time.sleep(0.5)


# ============================================================
# キー操作ヘルパー
# ============================================================

def tab(n=1):
    """Tabキーをn回押す"""
    for _ in range(n):
        pyautogui.press('tab')
        time.sleep(DELAY_KEY)


def input_field(text):
    """クリップボード経由でテキスト入力（日本語・英数字両対応）"""
    pyperclip.copy(str(text))
    pyautogui.hotkey('ctrl', 'v')
    time.sleep(DELAY_KEY)


def clear_and_input(text):
    """フィールドを全選択してクリアしてから入力"""
    pyautogui.hotkey('ctrl', 'a')
    time.sleep(0.05)
    input_field(text)


# ============================================================
# データ変換
# ============================================================

def format_seisanbi(raw):
    """生産日をYYYY/MM/DD形式に変換（入力: 20260312 または datetime）"""
    if hasattr(raw, 'strftime'):
        return raw.strftime('%Y/%m/%d')
    s = str(int(raw)) if isinstance(raw, float) else str(raw)
    if len(s) == 8 and s.isdigit():
        return f"{s[:4]}/{s[4:6]}/{s[6:8]}"
    return s


def format_time(raw):
    """時間をHHMM → HH:MM形式に変換（例: 900 → 09:00）。未設定はNone"""
    if raw is None or str(raw).strip() in ('', '—', '-', 'None'):
        return None
    s = str(int(raw)) if isinstance(raw, float) else str(raw).strip()
    s = s.zfill(4)
    return f"{s[:2]}:{s[2:4]}"


# ============================================================
# Excel読み込み
# ============================================================

def load_excel(filepath):
    """Excel出力2を読み込んでレコードリストを返す（マッピング済みのみ）"""
    wb = openpyxl.load_workbook(filepath)

    # シート「生産実績2」を優先、なければアクティブシートを使用
    if '生産実績2' in wb.sheetnames:
        ws = wb['生産実績2']
    else:
        ws = wb.active

    headers = [cell.value for cell in ws[1]]
    print(f"読み込みシート: {ws.title}")
    print(f"ヘッダー: {headers}")

    records = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row[0]:  # 生産日が空なら終了
            break
        record = dict(zip(headers, row))

        # 変換済みのみ処理
        if record.get('マッピング状態') != '変換済み':
            print(f"  スキップ（未設定）: 品番={record.get('アプリ品番')}")
            continue

        # 生産数量が0またはNoneはスキップ
        qty = record.get('生産数量')
        if not qty or (isinstance(qty, (int, float)) and qty <= 0):
            print(f"  スキップ（生産数0）: {record.get('基幹品番')}")
            continue

        records.append(record)

    print(f"\n入力対象: {len(records)} 件")
    return records


# ============================================================
# 1レコード入力
# ============================================================

def input_one_record(record, hwnd):
    """1レコードを基幹システムに入力"""
    activate_window(hwnd)

    seisanbi   = format_seisanbi(record['生産日'])
    hinban     = str(record['基幹品番']).strip()
    kouteijun  = str(record['工順']).strip()
    seisansu   = str(int(record['生産数量']))
    start_time = format_time(record.get('開始時間'))
    end_time   = format_time(record.get('終了時間'))

    print(f"  生産日={seisanbi}  品番={hinban}  工順={kouteijun}  "
          f"生産数={seisansu}  {start_time}～{end_time}")

    # ---- 処理区分（1:登録）→ Tabで生産日へ ----
    tab(TABS_TO_SEISANBI)

    # ---- 生産日 ----
    clear_and_input(seisanbi)

    # ---- 品番 ----
    tab(TABS_TO_HINBAN)
    clear_and_input(hinban)
    pyautogui.press('enter')        # 品番確定（品名・工程情報を引く）
    time.sleep(DELAY_HINBAN)

    # ---- 工程順位 ----
    tab(TABS_AFTER_HINBAN)
    clear_and_input(kouteijun)

    # ---- 生産数(完成) ----
    tab(TABS_TO_SEISANSU)
    clear_and_input(seisansu)

    # ---- 加工時間1 開始 ----
    if start_time:
        tab(TABS_TO_JIKAN_START)
        clear_and_input(start_time)

        # ---- 加工時間1 終了 ----
        if end_time:
            tab(TABS_TO_JIKAN_END)
            clear_and_input(end_time)

    # ---- F12 登録 ----
    pyautogui.press('f12')
    time.sleep(DELAY_REGISTER)


# ============================================================
# メイン
# ============================================================

def main():
    if len(sys.argv) < 2:
        print("使用方法: python kikan_input.py <Excelファイルパス>")
        print("例:       python kikan_input.py production_record2_20260310_20260312.xlsx")
        sys.exit(1)

    filepath = sys.argv[1]
    if not Path(filepath).exists():
        print(f"ファイルが見つかりません: {filepath}")
        sys.exit(1)

    # --- レコード読み込み ---
    records = load_excel(filepath)
    if not records:
        print("入力するレコードがありません")
        sys.exit(0)

    # --- 基幹システムのウィンドウ確認 ---
    hwnd = find_window(WINDOW_TITLE_PART)
    if not hwnd:
        print(f"\n基幹システムのウィンドウが見つかりません（検索: '{WINDOW_TITLE_PART}'）")
        print("SSE0040を開いてから再実行してください")
        sys.exit(1)

    title = win32gui.GetWindowText(hwnd)
    print(f"\n基幹システム検出: {title}")
    print(f"\n{int(DELAY_START)}秒後に入力を開始します。")
    print("★ 基幹システムの入力画面（SSE0040）で処理区分フィールドを選択した状態にしてください")
    print("★ 緊急停止: マウスを画面の左上角に移動")

    for i in range(int(DELAY_START), 0, -1):
        print(f"  {i}...")
        time.sleep(1)

    print("\n===== 入力開始 =====")
    pyautogui.FAILSAFE = True  # フェイルセーフ有効

    success = 0
    errors  = 0

    for i, record in enumerate(records):
        print(f"\n[{i+1}/{len(records)}] 入力中...")
        try:
            input_one_record(record, hwnd)
            success += 1
        except pyautogui.FailSafeException:
            print("\n緊急停止しました（フェイルセーフ）")
            break
        except Exception as e:
            print(f"  エラー: {e}")
            errors += 1
            ans = input("  続行しますか？ (y/n): ").strip().lower()
            if ans != 'y':
                print("  中断します")
                break

    print(f"\n===== 完了 =====")
    print(f"成功: {success}件  エラー: {errors}件")


if __name__ == "__main__":
    main()
