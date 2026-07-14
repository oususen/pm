from __future__ import annotations

import argparse
import re
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path

import fitz
from openpyxl import Workbook
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


NUM_RE = re.compile(r"^[+-]?\d+$")
DATE_RE = re.compile(r"^(?:\d{1,2}/\d{1,2}|\d{1,2}日)$")
MONTH_LABEL_RE = re.compile(r"^\d{1,2}月$")
CODE_RE_ORDER = re.compile(r"^[0-9A-Za-z][0-9A-Za-z\-]{5,}$")
CODE_RE_PROGRESS = re.compile(r"^[0-9A-Za-z][0-9A-Za-z\-]*[GＧ]$")


@dataclass
class ExtractedItem:
    code: str
    name: str
    page: int
    rows: dict[str, dict[str, int]]


def _normalize_code(code: str) -> str:
    text = str(code or "").strip().upper().replace("Ｇ", "G")
    return text[:-1] if text.endswith("G") else text


def _norm_name(text: str) -> str:
    return re.sub(r"\s+", "", str(text or "").replace("　", "").strip())


def _date_sort_key(label: str):
    text = str(label or "").strip()
    if not text:
        return (999, 999)
    if "/" in text:
        month, day = text.split("/", 1)
        try:
            return (int(month), int(day))
        except ValueError:
            return (999, 999)
    try:
        return (999, int(text.replace("日", "")))
    except ValueError:
        return (999, 999)


def _to_int(value) -> int:
    if value in (None, ""):
        return 0
    if isinstance(value, bool):
        return int(value)
    if isinstance(value, (int, float)):
        return int(value)
    text = str(value).strip()
    if not text:
        return 0
    try:
        return int(float(text))
    except ValueError:
        return 0


def _extract_progress_excel_items(excel_path: str) -> dict[str, ExtractedItem]:
    wb = load_workbook(excel_path, data_only=True)
    ws = wb["進度表"] if "進度表" in wb.sheetnames else wb[wb.sheetnames[0]]
    items: dict[str, ExtractedItem] = {}

    row_idx = 1
    while row_idx <= ws.max_row:
        code = str(ws.cell(row_idx, 1).value or "").strip()
        if not CODE_RE_PROGRESS.fullmatch(code):
            row_idx += 1
            continue

        name = str(ws.cell(row_idx + 1, 1).value or "").strip()
        header_row = row_idx + 1
        sub_header_row = row_idx + 2

        date_cols: list[tuple[int, str]] = []
        for col_idx in range(3, ws.max_column + 1):
            header_text = str(ws.cell(header_row, col_idx).value or "").strip()
            sub_text = str(ws.cell(sub_header_row, col_idx).value or "").strip()
            if DATE_RE.fullmatch(header_text) and sub_text:
                date_cols.append((col_idx, header_text))

        rows: dict[str, dict[str, int]] = {}
        data_start_row = sub_header_row + 1
        for offset in range(6):
            label = str(ws.cell(data_start_row + offset, 2).value or "").strip()
            if label not in {"内示", "確定"}:
                continue
            row: dict[str, int] = defaultdict(int)
            for col_idx, date_label in date_cols:
                value = _to_int(ws.cell(data_start_row + offset, col_idx).value)
                if value:
                    row[date_label] += value
            rows[label] = dict(row)

        normalized_code = _normalize_code(code)
        items[normalized_code] = ExtractedItem(
            code=code,
            name=name,
            page=1,
            rows=rows,
        )
        row_idx = data_start_row + 7

    return items


def _extract_items(pdf_path: str, kind: str) -> dict[str, ExtractedItem]:
    doc = fitz.open(pdf_path)
    items: dict[str, ExtractedItem] = {}

    for page_index in range(doc.page_count):
        page = doc.load_page(page_index)
        words = page.get_text("words")
        words.sort(key=lambda w: (w[1], w[0]))

        if kind == "progress":
            code_re = CODE_RE_PROGRESS
            row_labels = ["内示", "確定"]
            code_min_y = 20
        else:
            code_re = CODE_RE_ORDER
            row_labels = ["予定", "確定"]
            code_min_y = 80

        candidate_codes = [
            w for w in words
            if code_re.fullmatch(str(w[4]).strip())
            and w[1] > code_min_y
            and (
                # 注文書は左端の「品番」列のみ。右側の部品番号や数量を品番扱いしない。
                float(w[0]) < 90 if kind == "order" else float(w[0]) < 55
            )
        ]
        candidate_codes.sort(key=lambda w: (w[1], w[0]))

        for idx, code_word in enumerate(candidate_codes):
            code = str(code_word[4]).strip()
            code_y = float(code_word[1])
            code_x = float(code_word[0])
            next_y = float(candidate_codes[idx + 1][1]) if idx + 1 < len(candidate_codes) else float(page.rect.height)

            block_words = [
                w for w in words
                if code_y - 1 <= w[1] < next_y - 1
            ]

            # 日付ヘッダがあるブロックのみ明細として扱う
            date_header_words = [
                w for w in block_words
                if code_y + 4 <= w[1] <= code_y + 24
                and float(w[0]) > 80
                and DATE_RE.fullmatch(str(w[4]).strip())
            ]
            if not date_header_words:
                continue

            date_header_words.sort(key=lambda w: float(w[0]))
            date_centers = [((float(w[0]) + float(w[2])) / 2.0, str(w[4]).strip()) for w in date_header_words]
            min_date_x = min(center for center, _ in date_centers)

            if kind == "order":
                name_min_x, name_max_x = 90, 230
            else:
                name_min_x, name_max_x = 10, 105
            if kind == "order":
                name_y_min, name_y_max = code_y - 2, code_y + 12
            else:
                name_y_min, name_y_max = code_y + 5, code_y + 22
            name_words = [
                str(w[4]).strip()
                for w in block_words
                if name_y_min <= w[1] <= name_y_max
                and name_min_x <= float(w[0]) <= name_max_x
                and str(w[4]).strip() != code
                and not NUM_RE.fullmatch(str(w[4]).strip())
                and not DATE_RE.fullmatch(str(w[4]).strip())
                and not MONTH_LABEL_RE.fullmatch(str(w[4]).strip())
                and str(w[4]).strip() not in {
                    "繰越", "予定", "確定", "内示", "計画", "実績", "進度", "計進", "納入", "増減", "合計数"
                }
            ]
            name = " ".join(name_words).strip()
            if not name:
                continue

            rows: dict[str, dict[str, int]] = {}
            for label in row_labels:
                label_word = next(
                    (
                        w for w in block_words
                        if code_y + 15 <= w[1] <= code_y + 42 and str(w[4]).strip() == label
                    ),
                    None,
                )
                if not label_word:
                    continue

                label_y = float(label_word[1])
                cell_words = [
                    w for w in block_words
                    if abs(float(w[1]) - label_y) < 4
                    and float(w[0]) > min_date_x - 5
                    and NUM_RE.fullmatch(str(w[4]).strip())
                ]
                row: dict[str, int] = defaultdict(int)
                for cell_word in cell_words:
                    cell_x = (float(cell_word[0]) + float(cell_word[2])) / 2.0
                    nearest_index = min(
                        range(len(date_centers)),
                        key=lambda i: abs(date_centers[i][0] - cell_x),
                    )
                    # セル位置が大きくずれている誤検出は除外
                    if abs(date_centers[nearest_index][0] - cell_x) > 12:
                        continue
                    row[date_centers[nearest_index][1]] += int(str(cell_word[4]).strip())
                rows[label] = dict(row)

            normalized_code = _normalize_code(code)
            current = items.get(normalized_code)
            item = ExtractedItem(
                code=code,
                name=name,
                page=page_index + 1,
                rows=rows,
            )
            if current is None:
                items[normalized_code] = item
            else:
                # 先に出たものを優先しつつ、空でない行があれば補完する
                for key, row in item.rows.items():
                    if key not in current.rows or not current.rows[key]:
                        current.rows[key] = row

    return items


def _build_workbook(order_items: dict[str, ExtractedItem], progress_items: dict[str, ExtractedItem]) -> Workbook:
    wb = Workbook()
    thin = Side(style="thin", color="999999")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    sub_fill = PatternFill("solid", fgColor="D9EAF7")
    warn_fill = PatternFill("solid", fgColor="FFF2CC")
    diff_fill = PatternFill("solid", fgColor="FCE4D6")
    group_fills = [
        PatternFill("solid", fgColor="D9EAF7"),
        PatternFill("solid", fgColor="DDEBF7"),
        PatternFill("solid", fgColor="E2F0D9"),
        PatternFill("solid", fgColor="FCE4D6"),
        PatternFill("solid", fgColor="E4DFEC"),
    ]
    red_font = Font(color="C00000")
    blue_font = Font(color="0000FF")

    # シート1: 品番違い
    ws1 = wb.active
    ws1.title = "品番違い"
    ws1.append(["区分", "Gあり品番", "Gなし品番", "品名", "出典", "ページ", "内示合計", "確定合計", "需要有無"])
    for cell in ws1[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border
        cell.alignment = Alignment(horizontal="center")

    order_codes = set(order_items.keys())
    progress_codes = set(progress_items.keys())
    only_order = sorted(order_codes - progress_codes)
    only_progress = sorted(progress_codes - order_codes)

    for code in only_order:
        item = order_items[code]
        forecast_total = sum(item.rows.get("予定", {}).values())
        firm_total = sum(item.rows.get("確定", {}).values())
        demand_status = "あり" if forecast_total or firm_total else "なし"
        ws1.append(["注文書のみ", "", _normalize_code(item.code), item.name, "外作注文書", item.page, forecast_total, firm_total, demand_status])
    for code in only_progress:
        item = progress_items[code]
        forecast_total = sum(item.rows.get("内示", {}).values())
        firm_total = sum(item.rows.get("確定", {}).values())
        demand_status = "あり" if forecast_total or firm_total else "なし"
        ws1.append(["進度表のみ", item.code, _normalize_code(item.code), item.name, "進度表", item.page, forecast_total, firm_total, demand_status])

    for row in ws1.iter_rows(min_row=2):
        for cell in row:
            cell.border = border
        if row[0].value == "注文書のみ":
            for cell in row:
                cell.fill = warn_fill
        elif row[0].value == "進度表のみ":
            for cell in row:
                cell.fill = sub_fill

    ws1.freeze_panes = "A2"
    ws1.column_dimensions["A"].width = 14
    ws1.column_dimensions["B"].width = 16
    ws1.column_dimensions["C"].width = 16
    ws1.column_dimensions["D"].width = 28
    ws1.column_dimensions["E"].width = 14
    ws1.column_dimensions["F"].width = 8
    ws1.column_dimensions["G"].width = 10
    ws1.column_dimensions["H"].width = 10
    ws1.column_dimensions["I"].width = 10

    # シート2: 日付別数量差分
    ws2 = wb.create_sheet("数量差分")
    ws2.append(["日付", "Gあり品番", "Gなし品番", "品名", "区分", "注文書数量", "進度表数量", "数量差"])
    for cell in ws2[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border
        cell.alignment = Alignment(horizontal="center")

    row_map = [("予定", "内示"), ("確定", "確定")]
    common_codes = sorted(order_codes & progress_codes)
    for code in common_codes:
        order_item = order_items[code]
        progress_item = progress_items[code]
        for order_label, progress_label in row_map:
            order_row = order_item.rows.get(order_label, {})
            progress_row = progress_item.rows.get(progress_label, {})
            all_dates = sorted(set(order_row.keys()) | set(progress_row.keys()), key=_date_sort_key)
            for target_date in all_dates:
                order_qty = int(order_row.get(target_date, 0) or 0)
                progress_qty = int(progress_row.get(target_date, 0) or 0)
                diff = order_qty - progress_qty
                if diff == 0:
                    continue
                ws2.append([
                    target_date,
                    progress_item.code,
                    _normalize_code(progress_item.code),
                    order_item.name,
                    "内示" if progress_label == "内示" else "確定",
                    order_qty,
                    progress_qty,
                    diff,
                ])

    last_code = None
    group_index = -1
    for row in ws2.iter_rows(min_row=2):
        current_code = row[2].value
        if current_code != last_code:
            group_index += 1
            last_code = current_code
        group_fill = group_fills[group_index % len(group_fills)]
        for cell in row:
            cell.border = border
            cell.fill = group_fill
        if row[6].value is not None and row[6].value != 0:
            row[6].font = red_font if row[6].value > 0 else blue_font
            if row[6].value > 0:
                row[7].fill = diff_fill

    ws2.freeze_panes = "A2"
    ws2.column_dimensions["A"].width = 12
    ws2.column_dimensions["B"].width = 16
    ws2.column_dimensions["C"].width = 16
    ws2.column_dimensions["D"].width = 28
    ws2.column_dimensions["E"].width = 10
    ws2.column_dimensions["F"].width = 12
    ws2.column_dimensions["G"].width = 12
    ws2.column_dimensions["H"].width = 10

    # シート3: その他差分
    ws3 = wb.create_sheet("その他差分")
    ws3.append(["品番", "注文書品名", "進度表品名", "差分種別", "注文書ページ", "進度表ページ", "備考"])
    for cell in ws3[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border
        cell.alignment = Alignment(horizontal="center")

    for code in common_codes:
        order_item = order_items[code]
        progress_item = progress_items[code]
        order_name = _norm_name(order_item.name)
        progress_name = _norm_name(progress_item.name)
        if order_name != progress_name:
            ws3.append([
                order_item.code,
                order_item.name,
                progress_item.name,
                "品名差分",
                order_item.page,
                progress_item.page,
                "",
            ])

        # 抽出できたのに日付別の行が空のケースは注意として残す
        if not order_item.rows or not progress_item.rows:
            ws3.append([
                order_item.code,
                order_item.name,
                progress_item.name,
                "明細抽出不足",
                order_item.page,
                progress_item.page,
                "日付別セルの抽出に不足あり",
            ])

    for row in ws3.iter_rows(min_row=2):
        for cell in row:
            cell.border = border
    ws3.freeze_panes = "A2"
    ws3.column_dimensions["A"].width = 16
    ws3.column_dimensions["B"].width = 28
    ws3.column_dimensions["C"].width = 28
    ws3.column_dimensions["D"].width = 14
    ws3.column_dimensions["E"].width = 10
    ws3.column_dimensions["F"].width = 10
    ws3.column_dimensions["G"].width = 30

    return wb


def _parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="外作注文書PDFと進度表PDFを突合せ、品番差分・数量差分・その他差分をExcel出力します。",
    )
    parser.add_argument(
        "--order-pdf",
        default="外作注文書000095_20260605084453.pdf",
        help="外作注文書PDFのパス",
    )
    parser.add_argument(
        "--progress-pdf",
        default="進度表_000095_2026-06-05 (3).pdf",
        help="進度表PDFのパス",
    )
    parser.add_argument(
        "--output",
        default="比較表_外作注文書_進度表_000095_需要確認_G列分け.xlsx",
        help="出力Excelのパス",
    )
    return parser.parse_args()


def main() -> None:
    args = _parse_args()
    order_pdf = Path(args.order_pdf)
    progress_pdf = Path(args.progress_pdf)
    output_xlsx = Path(args.output)

    order_items = _extract_items(str(order_pdf), "order")
    progress_items = _extract_items(str(progress_pdf), "progress")

    wb = _build_workbook(order_items, progress_items)
    wb.save(output_xlsx)
    print(f"saved: {output_xlsx}")
    print(f"order items: {len(order_items)}")
    print(f"progress items: {len(progress_items)}")


if __name__ == "__main__":
    main()
