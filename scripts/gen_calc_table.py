"""
在庫・進度・計画在庫・計画進度 計算ロジック一覧表 Excel生成スクリプト
"""
from openpyxl import Workbook
from openpyxl.styles import (
    Alignment, Border, Font, PatternFill, Side
)
from openpyxl.utils import get_column_letter

wb = Workbook()
ws = wb.active
ws.title = "計算ロジック一覧"

# ---- スタイル定義 ----
HEADER_ROW_FILL  = PatternFill("solid", fgColor="1F4E79")
HEADER_COL_FILL  = PatternFill("solid", fgColor="2E75B6")
SECTION_FILL     = PatternFill("solid", fgColor="D6E4F0")
ALT_FILL         = PatternFill("solid", fgColor="EBF3FB")
WHITE_FILL       = PatternFill("solid", fgColor="FFFFFF")

HEADER_FONT      = Font(name="Meiryo UI", bold=True, color="FFFFFF", size=10)
LABEL_FONT       = Font(name="Meiryo UI", bold=True, color="1F4E79", size=9)
BODY_FONT        = Font(name="Meiryo UI", size=9)
TITLE_FONT       = Font(name="Meiryo UI", bold=True, color="1F4E79", size=13)

thin = Side(style="thin", color="B0C4DE")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

def cell(ws, row, col, value="", fill=None, font=None, align=None, border=None):
    c = ws.cell(row=row, column=col, value=value)
    if fill:   c.fill   = fill
    if font:   c.font   = font
    if align:  c.alignment = align
    if border: c.border = border
    return c

# ---- タイトル ----
ws.merge_cells("A1:F1")
cell(ws, 1, 1, "在庫・計画在庫・進度・計画進度　計算ロジック一覧表",
     fill=PatternFill("solid", fgColor="1F4E79"),
     font=Font(name="Meiryo UI", bold=True, color="FFFFFF", size=14),
     align=CENTER)
ws.row_dimensions[1].height = 28

ws.merge_cells("A2:F2")
cell(ws, 2, 1, "※ 日替わり時刻 = 8:00  |  calc_start_date = today − (max親BOM LT + 1) 営業日",
     fill=PatternFill("solid", fgColor="D6E4F0"),
     font=Font(name="Meiryo UI", italic=True, color="1F4E79", size=9),
     align=Alignment(horizontal="left", vertical="center"))
ws.row_dimensions[2].height = 16

# ---- ヘッダ行 ----
headers = ["項目", "在庫\n(stock_qty)", "計画在庫\n(planned_stock_qty)", "進度\n(progress_qty)", "計画進度\n(planned_progress_qty)", "備考"]
ROW_H = 3
for col, h in enumerate(headers, 1):
    cell(ws, ROW_H, col, h, fill=HEADER_ROW_FILL, font=HEADER_FONT, align=CENTER, border=BORDER)
ws.row_dimensions[ROW_H].height = 30

# ---- データ定義 ----
rows = [
    ("関数名",
     "recalculate_stock_qty()",
     "recalculate_planned_stock_qty()",
     "recalculate_progress_qty()",
     "recalculate_progress_qty()\n※同一関数内で両方計算",
     ""),
    ("ファイル / 行番号",
     "inventory_calculator.py\nL772",
     "inventory_calculator.py\nL1015",
     "progress_calculator.py\nL17",
     "progress_calculator.py\nL17",
     ""),
    ("計算式",
     "前日在庫\n＋ 実績\n－ 出荷実績\n＋ 仕損調整\n＋ 調整",
     "前日計画在庫\n＋ 実績(過去) / 計画(未来)\n－ 計画出荷\n＋ 調整",
     "前日進度\n＋ 実績\n－ 需要\n＋ 調整\n－ 仕損",
     "前日計画進度\n＋ 計画\n－ 需要\n＋ 調整\n－ 仕損",
     ""),
    ("計算開始日（内部）\neffective_start",
     "min(start_date, calc_start_date)\n\ncalc_start_date\n= today − (max親BOM LT + 1) 営業日\n\n※ 計画在庫の初期値参照日と\n　 同じ下限まで再計算",
     "calc_start_date\n= today − (max親BOM LT + 1) 営業日",
     "calc_start_date\n= today − (max親BOM LT + 1) 営業日",
     "同左",
     "在庫と計画在庫の\n開始日を統一するため\n在庫側を拡張（修正済）"),
    ("更新対象日範囲",
     "effective_start 〜 end_date\n（effective_start 以前は更新しない）",
     "calc_start_date 〜 end_date",
     "calc_start_date 〜 end_date",
     "同左",
     ""),
    ("初期値の取得",
     "effective_start より前の\n最新 stock_qty を DB から取得",
     "【最終品】\ncalc_start_date の stock_qty\n− LT日分の firm 需要\n\n【中間品】\ncalc_start_date 前日の\nplanned_stock_qty",
     "calc_start_date 前日の\nprogress_qty を DB から取得",
     "calc_start_date 前日の\nprogress_qty（実進度）を DB から取得\n※ planned_progress_qty ではない\n　 計画在庫が実在庫を初期値にするのと同じ思想",
     ""),
    ("過去日（today 未満）",
     "actual_qty を使用",
     "actual_qty を使用",
     "actual_qty を使用",
     "actual_qty を使用\n（計画在庫と同じ時制考慮）",
     ""),
    ("未来日（today 以降）",
     "actual_qty を使用\n（実績ベース）",
     "plan_qty を使用",
     "actual_qty を使用",
     "plan_qty を使用",
     ""),
    ("需要データソース",
     "OrderLine\n（firm 受注, LT シフト済）",
     "OrderLine（firm）\n↓ なければ order_qty にフォールバック",
     "LineDemand\n（firm_qty / forecast_qty）\n※ LineBacklog 不可",
     "LineDemand\n（firm_qty / forecast_qty）\n※ LineBacklog 不可",
     "進度は LineDemand\n必須（業務ルール）"),
    ("中間品の出荷計算",
     "親の actual_qty\n＋ scrap_qty\n× BOM 数量",
     "【過去】親の actual or plan\n【未来】親の plan_qty\n　　　＋ scrap_qty × BOM 数量",
     "−",
     "−",
     ""),
    ("手動調整テーブル",
     "LineBacklogAdjustment\n(STOCK)",
     "LineBacklogAdjustment\n(PLANNED_STOCK)",
     "LineBacklogAdjustment\n(PROGRESS)",
     "LineBacklogAdjustment\n(PLANNED_PROGRESS)",
     ""),
    ("仕損の扱い",
     "scrap_adjust_qty\n（他工程仕損, マイナス値）\nを加算",
     "同左",
     "scrap_adjust_qty\n（他工程仕損）\nを減算",
     "同左",
     "自工程仕損は scrap_qty\nに格納"),
    ("非営業日",
     "前営業日の値を引き継ぎ\n（更新なし）",
     "同左",
     "同左",
     "同左",
     ""),
    ("sequence_no=0 行",
     "各日付の代表行\n（需要・在庫値を保持）",
     "同左",
     "進度値を格納\n（plan_qty = 0 必須）",
     "同左",
     "sequence_no > 0 は\n計画値専用行"),
]

for i, row_data in enumerate(rows):
    r = i + ROW_H + 1
    fill = ALT_FILL if i % 2 == 0 else WHITE_FILL
    for col, val in enumerate(row_data, 1):
        f = LABEL_FONT if col == 1 else BODY_FONT
        bg = HEADER_COL_FILL if col == 1 else fill
        ft = Font(name="Meiryo UI", bold=True, color="FFFFFF", size=9) if col == 1 else f
        cell(ws, r, col, val, fill=bg, font=ft, align=WRAP, border=BORDER)
    ws.row_dimensions[r].height = 52

# ---- 列幅 ----
col_widths = [22, 30, 38, 32, 36, 26]
for col, w in enumerate(col_widths, 1):
    ws.column_dimensions[get_column_letter(col)].width = w

# ---- セクション2: 呼び出しエントリポイント ----
sep_row = ROW_H + len(rows) + 2

ws.merge_cells(f"A{sep_row}:F{sep_row}")
cell(ws, sep_row, 1, "呼び出しエントリポイント",
     fill=PatternFill("solid", fgColor="1F4E79"),
     font=Font(name="Meiryo UI", bold=True, color="FFFFFF", size=11),
     align=CENTER)
ws.row_dimensions[sep_row].height = 22

ep_headers = ["項目", "内容", "", "", "", ""]
for col, h in enumerate(ep_headers[:2], 1):
    cell(ws, sep_row + 1, col, h if h else "",
         fill=HEADER_ROW_FILL, font=HEADER_FONT, align=CENTER, border=BORDER)
ws.merge_cells(f"B{sep_row+1}:F{sep_row+1}")
cell(ws, sep_row + 1, 2, "内容",
     fill=HEADER_ROW_FILL, font=HEADER_FONT, align=CENTER, border=BORDER)
ws.row_dimensions[sep_row + 1].height = 20

ep_rows = [
    ("統合関数", "recalculate_inventory_for_line(\n  line_id, start_date, end_date,\n  include_progress=True,\n  progress_only=False,\n  line_final_only=False,\n  product_ids=None\n)"),
    ("progress_only モード", "True のとき:\n  仕損集計・firm_map 構築をスキップ\n  進度・計画進度のみ計算（在庫・計画在庫はスキップ）\n\n使用箇所: 在庫進度調整画面のPROGRESS/PLANNED_PROGRESS一括再計算"),
    ("実行順（通常）", "① 仕損集計 (aggregate_scrap_to_backlog)\n② firm_map 構築（確定受注 × LT シフト）\n③ 調整マップ構築（STOCK / PLANNED_STOCK / PROGRESS / PLANNED_PROGRESS）\n④ 製品ごとに 在庫 → 計画在庫 → 進度・計画進度 の順に計算して DB へ一括更新"),
    ("通常トリガー", "日次スケジューラ  scheduler/tasks.py"),
    ("手動トリガー", "実績入力・計画変更・受注変更時（各 View / Signal から呼び出し）"),
    ("棚卸時", "stocktake_initializer.py の専用関数を使用\n※ inventory_calculator.py / progress_calculator.py は変更しない（業務ルール）"),
]

for i, (label, content) in enumerate(ep_rows):
    r = sep_row + 2 + i
    fill = ALT_FILL if i % 2 == 0 else WHITE_FILL
    cell(ws, r, 1, label,
         fill=HEADER_COL_FILL,
         font=Font(name="Meiryo UI", bold=True, color="FFFFFF", size=9),
         align=WRAP, border=BORDER)
    ws.merge_cells(f"B{r}:F{r}")
    cell(ws, r, 2, content, fill=fill, font=BODY_FONT, align=WRAP, border=BORDER)
    ws.row_dimensions[r].height = 48

# ---- ウィンドウ枠固定 ----
ws.freeze_panes = "B4"

out = r"d:\pm\仕様書\在庫進度計算ロジック一覧.xlsx"
wb.save(out)
print(f"保存完了: {out}")
